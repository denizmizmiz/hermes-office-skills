#!/usr/bin/env python3
# MIT License. Part of the Hermes xlsx skill.
"""Create an .xlsx workbook from a JSON spec.

Spec (JSON object):
  {
    "full_calc_on_load": true,          # force recalculation on open (optional)
    "defined_names": {"Rates": "'Data'!$B$2:$B$4"},   # workbook scope
    "sheets": [
      {
        "name": "Data",
        "rows": [["Header", 1, true], ...],   # scalars or cell objects (below)
        "cells": {"A1": {"value": 5, "format": "0.00%"}},  # sparse overrides
        "column_widths": {"A": 22, "B": 12},
        "row_heights": {"1": 24},
        "merges": ["A1:C1"],
        "freeze_panes": "A2",
        "autofilter": "A1:C10",
        "conditional_formats": [
          {"range": "B2:B9", "type": "cell_is", "operator": "greaterThan",
           "formula": ["100"], "fill": "FFC7CE"},
          {"range": "C2:C9", "type": "color_scale"}
        ],
        "charts": [
          {"type": "bar", "title": "Sales", "anchor": "F2",
           "data": "B1:B5", "categories": "A2:A5"}
        ],
        "validations": [
          {"range": "D2:D9", "type": "list", "formula1": "\"Yes,No,Maybe\""}
        ],
        "tables": [
          {"name": "Sales", "range": "A1:C4",
           "style": "TableStyleMedium9"}        # native Excel table
        ],
        "protection": {"password": "your-password",   # NOT security --
                       "unlock": ["B2:B9"]}           # see SKILL.md Pitfalls
      }
    ]
  }

Cell object keys (all optional except value/formula):
  value          scalar; JSON true/false -> bool, numbers stay numeric
  type           "date" or "datetime" -> value parsed from ISO string
  formula        e.g. "=SUM(A2:A9)" (leading '=' optional)
  hyperlink      URL; value becomes the display text
  note           cell note text (or {"text": ..., "author": ...})
  format         Excel number format, e.g. "$#,##0.00", "0.0%", "yyyy-mm-dd"
  bold, italic   booleans
  font_size      points
  font_color     hex RGB like "FF0000"
  fill           solid fill hex RGB like "DDEBF7"
  border         "thin" | "medium" | "thick" (all four sides)
  align          "left" | "center" | "right"
  valign         "top" | "center" | "bottom"
  wrap           boolean (wrap text)

Validation and safety (Phase 1):
  The spec is validated BEFORE anything is built. Missing required keys and
  type mismatches fail with SPEC_INVALID and a 'path' pointing at the exact
  field; unknown keys fail with SPEC_UNKNOWN_KEY unless --allow-unknown-keys
  is passed (then they are returned as warnings). The output goes through the
  shared safe-write chain (backup -> temp -> validate -> atomic commit); an
  existing file at the output path is backed up first, never silently
  destroyed.

Usage:
  xlsx_create.py spec.json out.xlsx
  xlsx_create.py - out.xlsx   (spec on stdin)
  xlsx_create.py spec.json out.xlsx --allow-unknown-keys

Prints a JSON summary to stdout; exits non-zero on failure.
"""
from __future__ import annotations

import argparse
import json
import sys
from datetime import date, datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

import xlsx_common as common  # noqa: E402

try:  # dependency guard: structured DEPENDENCY_MISSING, never a traceback
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, LineChart, PieChart, Reference
    from openpyxl.comments import Comment
    from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
    from openpyxl.styles import (Alignment, Border, Font, PatternFill,
                                 Protection, Side)
    from openpyxl.utils import column_index_from_string, range_boundaries
    from openpyxl.workbook.defined_name import DefinedName
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.worksheet.table import Table, TableStyleInfo
except ImportError as _exc:  # pragma: no cover - environment guard
    sys.exit(common.dependency_error(_exc))


def parse_typed(value, type_hint=None):
    if type_hint == "date" and isinstance(value, str):
        return date.fromisoformat(value)
    if type_hint == "datetime" and isinstance(value, str):
        return datetime.fromisoformat(value)
    return value


def apply_cell(ws, coord, spec):
    cell = ws[coord]
    if isinstance(spec, dict):
        if "formula" in spec:
            f = spec["formula"]
            cell.value = f if f.startswith("=") else "=" + f
        elif "value" in spec:
            cell.value = parse_typed(spec["value"], spec.get("type"))
        if "hyperlink" in spec:
            cell.hyperlink = spec["hyperlink"]
            if cell.value is None:
                cell.value = spec["hyperlink"]
            cell.style = "Hyperlink"
        if "note" in spec:
            note = spec["note"]
            if isinstance(note, dict):
                cell.comment = Comment(note.get("text", ""),
                                       note.get("author", "xlsx-skill"))
            else:
                cell.comment = Comment(str(note), "xlsx-skill")
        if "format" in spec:
            cell.number_format = spec["format"]
        font_kw = {}
        if spec.get("bold"):
            font_kw["bold"] = True
        if spec.get("italic"):
            font_kw["italic"] = True
        if "font_size" in spec:
            font_kw["size"] = spec["font_size"]
        if "font_color" in spec:
            font_kw["color"] = spec["font_color"]
        if font_kw:
            cell.font = Font(**font_kw)
        if "fill" in spec:
            cell.fill = PatternFill("solid", fgColor=spec["fill"])
        if "border" in spec:
            side = Side(style=spec["border"])
            cell.border = Border(left=side, right=side, top=side, bottom=side)
        align_kw = {}
        if "align" in spec:
            align_kw["horizontal"] = spec["align"]
        if "valign" in spec:
            align_kw["vertical"] = spec["valign"]
        if spec.get("wrap"):
            align_kw["wrap_text"] = True
        if align_kw:
            cell.alignment = Alignment(**align_kw)
    else:
        cell.value = spec


def ref_from_range(ws, rng):
    min_col, min_row, max_col, max_row = range_boundaries(rng)
    return Reference(ws, min_col=min_col, min_row=min_row,
                     max_col=max_col, max_row=max_row)


def add_chart(ws, spec):
    kind = spec.get("type", "bar")
    chart = {"bar": BarChart, "line": LineChart, "pie": PieChart}[kind]()
    if "title" in spec:
        chart.title = spec["title"]
    data = ref_from_range(ws, spec["data"])
    chart.add_data(data, titles_from_data=spec.get("titles_from_data", True))
    if "categories" in spec:
        chart.set_categories(ref_from_range(ws, spec["categories"]))
    ws.add_chart(chart, spec.get("anchor", "H2"))


def add_conditional(ws, spec):
    rng = spec["range"]
    kind = spec.get("type", "cell_is")
    if kind == "color_scale":
        rule = ColorScaleRule(
            start_type="min", start_color=spec.get("start_color", "FFF8696B"),
            end_type="max", end_color=spec.get("end_color", "FF63BE7B"))
    else:
        fill = PatternFill("solid", fgColor=spec.get("fill", "FFC7CE"))
        rule = CellIsRule(operator=spec.get("operator", "greaterThan"),
                          formula=spec.get("formula", ["0"]), fill=fill)
    ws.conditional_formatting.add(rng, rule)


def build_sheet(ws, spec):
    for row in spec.get("rows", []):
        values, styled = [], []
        for item in row:
            if isinstance(item, dict):
                values.append(None)
                styled.append(item)
            else:
                values.append(item)
                styled.append(None)
        ws.append(values)
        r = ws.max_row
        for idx, item in enumerate(styled, start=1):
            if item is not None:
                apply_cell(ws, ws.cell(row=r, column=idx).coordinate, item)
    for coord, cell_spec in spec.get("cells", {}).items():
        apply_cell(ws, coord, cell_spec)
    for col, width in spec.get("column_widths", {}).items():
        ws.column_dimensions[col].width = width
    for row, height in spec.get("row_heights", {}).items():
        ws.row_dimensions[int(row)].height = height
    for rng in spec.get("merges", []):
        ws.merge_cells(rng)
    if spec.get("freeze_panes"):
        ws.freeze_panes = spec["freeze_panes"]
    if spec.get("autofilter"):
        ws.auto_filter.ref = spec["autofilter"]
    for cf in spec.get("conditional_formats", []):
        add_conditional(ws, cf)
    for ch in spec.get("charts", []):
        add_chart(ws, ch)
    for dv_spec in spec.get("validations", []):
        dv = DataValidation(type=dv_spec.get("type", "list"),
                            formula1=dv_spec["formula1"],
                            allow_blank=dv_spec.get("allow_blank", True))
        dv.add(dv_spec["range"])
        ws.add_data_validation(dv)
    for t_spec in spec.get("tables", []):
        table = Table(displayName=t_spec["name"], ref=t_spec["range"])
        table.tableStyleInfo = TableStyleInfo(
            name=t_spec.get("style", "TableStyleMedium9"),
            showRowStripes=t_spec.get("row_stripes", True),
            showColumnStripes=t_spec.get("column_stripes", False))
        ws.add_table(table)
    prot = spec.get("protection")
    if prot:
        for rng in prot.get("unlock", []):
            for row in ws[rng]:
                for cell in row:
                    cell.protection = Protection(locked=False)
        if prot.get("password"):
            ws.protection.password = prot["password"]
        ws.protection.sheet = True


# --------------------------------------------------------------------------
# spec schema  (D2/D3 -- mirrors the real spec shapes documented above)
# --------------------------------------------------------------------------

TABLE_SCHEMA = common.Schema(
    required=["name", "range"],
    optional=["style", "row_stripes", "column_stripes"],
    types={"name": "str", "range": "str", "style": "str",
           "row_stripes": "bool", "column_stripes": "bool"})

CHART_SCHEMA = common.Schema(
    required=["data"],
    optional=["type", "title", "anchor", "categories", "titles_from_data"],
    types={"type": "str", "title": "str", "anchor": "str", "data": "str",
           "categories": "str", "titles_from_data": "bool"})

CF_SCHEMA = common.Schema(
    required=["range"],
    optional=["type", "operator", "formula", "fill", "start_color",
              "end_color"],
    types={"range": "str", "type": "str", "operator": "str",
           "formula": "list", "fill": "str", "start_color": "str",
           "end_color": "str"})

VALIDATION_SCHEMA = common.Schema(
    required=["range", "formula1"],
    optional=["type", "allow_blank"],
    types={"range": "str", "formula1": "str", "type": "str",
           "allow_blank": "bool"})

PROTECTION_SCHEMA = common.Schema(
    optional=["password", "unlock"],
    types={"password": "str", "unlock": "list"})

SHEET_SCHEMA = common.Schema(
    optional=["name", "rows", "cells", "column_widths", "row_heights",
              "merges", "freeze_panes", "autofilter", "conditional_formats",
              "charts", "validations", "tables", "protection"],
    types={"name": "str", "rows": "list", "cells": "dict",
           "column_widths": "dict", "row_heights": "dict", "merges": "list",
           "freeze_panes": "str", "autofilter": "str",
           "conditional_formats": "list", "charts": "list",
           "validations": "list", "tables": "list", "protection": "dict"},
    children={"conditional_formats": CF_SCHEMA, "charts": CHART_SCHEMA,
              "validations": VALIDATION_SCHEMA, "tables": TABLE_SCHEMA,
              "protection": PROTECTION_SCHEMA},
    many=["conditional_formats", "charts", "validations", "tables"])

CREATE_SCHEMA = common.Schema(
    required=["sheets"],
    optional=["full_calc_on_load", "defined_names"],
    types={"sheets": "list", "full_calc_on_load": "bool",
           "defined_names": "dict"},
    children={"sheets": SHEET_SCHEMA},
    many=["sheets"])

CELL_KEYS = {"value", "type", "formula", "hyperlink", "note", "format",
             "bold", "italic", "font_size", "font_color", "fill", "border",
             "align", "valign", "wrap"}
CELL_TYPE_VALUES = {"date", "datetime"}

common.register_known_keys(
    sorted(CREATE_SCHEMA.known | SHEET_SCHEMA.known | CF_SCHEMA.known
           | CHART_SCHEMA.known | VALIDATION_SCHEMA.known
           | TABLE_SCHEMA.known | PROTECTION_SCHEMA.known | CELL_KEYS))


def check_cell_object(item, path, errors, unknown):
    """Cell objects (inside rows[][] or cells{}): unknown keys + type enum."""
    for key in sorted(item):
        if key not in CELL_KEYS:
            unknown.append({"path": f"{path}.{key}",
                            "value_type": type(item[key]).__name__})
    type_hint = item.get("type")
    if type_hint is not None and type_hint not in CELL_TYPE_VALUES:
        errors.append({"path": f"{path}.type",
                       "expected": "date | datetime",
                       "actual": repr(type_hint)})


def collect_extra_problems(spec):
    """Checks the generic Schema cannot express: cell objects and enums."""
    errors, unknown = [], []
    if not isinstance(spec, dict):
        return errors, unknown
    sheets = spec.get("sheets")
    if isinstance(sheets, list) and not sheets:
        errors.append({"path": "sheets", "expected": "at least one sheet",
                       "actual": "empty list"})
    if not isinstance(sheets, list):
        return errors, unknown
    for si, sheet in enumerate(sheets):
        if not isinstance(sheet, dict):
            continue
        spath = f"sheets[{si}]"
        rows = sheet.get("rows")
        if isinstance(rows, list):
            for ri, row in enumerate(rows):
                if not isinstance(row, list):
                    continue
                for ci, item in enumerate(row):
                    if isinstance(item, dict):
                        check_cell_object(item, f"{spath}.rows[{ri}][{ci}]",
                                          errors, unknown)
        cells = sheet.get("cells")
        if isinstance(cells, dict):
            for coord, item in cells.items():
                if isinstance(item, dict):
                    check_cell_object(item, f"{spath}.cells[{coord}]",
                                      errors, unknown)
        cfs = sheet.get("conditional_formats")
        if isinstance(cfs, list):
            for ci, item in enumerate(cfs):
                if not isinstance(item, dict):
                    continue
                kind = item.get("type")
                if kind is not None and kind not in ("cell_is",
                                                     "color_scale"):
                    errors.append({
                        "path": f"{spath}.conditional_formats[{ci}].type",
                        "expected": "cell_is | color_scale",
                        "actual": repr(kind)})
        charts = sheet.get("charts")
        if isinstance(charts, list):
            for ci, item in enumerate(charts):
                if not isinstance(item, dict):
                    continue
                kind = item.get("type")
                if kind is not None and kind not in ("bar", "line", "pie"):
                    errors.append({
                        "path": f"{spath}.charts[{ci}].type",
                        "expected": "bar | line | pie",
                        "actual": repr(kind)})
    return errors, unknown


def main(argv=None):
    ap = argparse.ArgumentParser(description="Create .xlsx from a JSON spec.")
    ap.add_argument("spec", help="path to JSON spec, or '-' for stdin")
    ap.add_argument("output", help="output .xlsx path")
    ap.add_argument("--allow-unknown-keys", action="store_true",
                    help="accept unknown spec keys with a warning instead "
                         "of failing (SPEC_UNKNOWN_KEY)")
    args = ap.parse_args(argv)

    try:
        if args.spec == "-":
            spec = json.load(sys.stdin)
        else:
            with open(args.spec, encoding="utf-8") as fh:
                spec = json.load(fh)
    except json.JSONDecodeError as exc:
        raise common.XlsxError(
            "SPEC_INVALID",
            f"Spec is not valid JSON: {exc}",
            recovery="Fix the JSON syntax and re-run.",
            context={"spec": args.spec}) from exc

    result = common.Result(mode="create", output=str(args.output))

    # D3: validate before building anything; strict on the write path.
    validation = common.validate_spec(spec, CREATE_SCHEMA, strict=True,
                                      allow_unknown=args.allow_unknown_keys)
    errors, unknown = collect_extra_problems(spec)
    validation["errors"].extend(errors)
    validation["unknown_keys"].extend(unknown)
    common.spec_error(result, validation,
                      strict=not args.allow_unknown_keys)

    wb = Workbook()
    wb.remove(wb.active)
    for sheet_spec in spec.get("sheets", []):
        ws = wb.create_sheet(sheet_spec.get("name", "Sheet1"))
        build_sheet(ws, sheet_spec)
    for name, ref in spec.get("defined_names", {}).items():
        wb.defined_names[name] = DefinedName(name, attr_text=ref)
    if spec.get("full_calc_on_load"):
        wb.calculation.fullCalcOnLoad = True

    # Safe write (F): the fresh workbook goes through the shared chain too;
    # an existing file at the output path is backed up before replacement.
    loaded = common.LoadedWorkbook(wb, path=None, mode="normal")
    save_info = common.save_workbook_safe(
        loaded, args.output, backup=True, expect_sheets=wb.sheetnames,
        plan={"action": "create", "sheets": list(wb.sheetnames)})

    result.set(sheets=list(wb.sheetnames),
               output=save_info["output"], backup=save_info["backup"],
               atomic=save_info["atomic"], bytes=save_info["bytes"],
               operation_id=save_info["operation_id"],
               plan_hash=save_info["plan_hash"],
               approval=save_info["approval"],
               manifest=save_info["manifest"])
    if save_info["manifest"] and not save_info["manifest"].get("ok"):
        result.warn("FEATURE_PARTIAL",
                    "the write manifest could not be updated: "
                    + str(save_info["manifest"].get("message")))
    if save_info["backup"]:
        result.diagnose(
            f"an existing file at the output path was backed up to "
            f"{save_info['backup']}")
    return result.emit()


if __name__ == "__main__":
    sys.exit(common.guard(main)())
