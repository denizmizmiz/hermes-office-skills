#!/usr/bin/env python3
# MIT License. Part of the Hermes xlsx skill.
"""Read an .xlsx workbook: inventory, JSON/CSV dumps, formula listing.

Modes (pick one):
  --sheets     JSON inventory: sheet names, dimensions, row/col counts
  --json       dump one sheet's rows as a JSON array of arrays
  --csv        dump one sheet as CSV to stdout or --out
  --formulas   JSON list of formula cells {"cell", "formula", "cached"}
  --notes      JSON list of cell notes/comments across sheets
  --names      JSON map of workbook defined names

Options:
  --sheet NAME     sheet to dump (default: active sheet)
  --data-only      load cached formula RESULTS instead of formula strings.
                   Caveat: openpyxl never computes formulas; cached values
                   exist only if the file was last saved by Excel/LibreOffice.
  --encoding ENC   encoding for --csv --out files (default utf-8)
  --out PATH       write --csv output to a file instead of stdout

Every JSON result carries warnings[]/unsupported[]/diagnostics[] (I3).
This script never writes to a workbook; the data_only handle it may open
is structurally non-writable (see xlsx_common, I2 guard). The --formulas
mode uses two separate read-only passes (formula text / cached values)
and never conflates them.

Usage:
  xlsx_read.py book.xlsx --sheets
  xlsx_read.py book.xlsx --json --sheet Data
  xlsx_read.py book.xlsx --csv --sheet Data --out data.csv
  xlsx_read.py book.xlsx --formulas
  xlsx_read.py book.xlsx --notes
  xlsx_read.py book.xlsx --names
"""
from __future__ import annotations

import argparse
import csv
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

import xlsx_common as common  # noqa: E402


def resolve_sheet(sheetnames, wanted, active):
    """Return the requested sheet name or raise a structured SHEET_NOT_FOUND."""
    if wanted is None:
        return active
    if wanted not in sheetnames:
        raise common.XlsxError(
            "SHEET_NOT_FOUND",
            f"Worksheet '{wanted}' does not exist.",
            recovery="Use --sheets to list the available sheet names.",
            context={"requested": wanted, "available": sheetnames},
        )
    return wanted


def cmd_sheets(wb, result):
    info = []
    for ws in wb.worksheets:
        info.append({
            "name": ws.title,
            "dimensions": ws.dimensions,
            "max_row": ws.max_row,
            "max_col": ws.max_column,
            "merged": [str(r) for r in ws.merged_cells.ranges],
            "charts": len(getattr(ws, "_charts", [])),
            "freeze_panes": ws.freeze_panes,
            "autofilter": ws.auto_filter.ref,
            "tables": {t.displayName: t.ref for t in ws.tables.values()},
            "protected": bool(ws.protection.sheet),
        })
    names = {name: dn.attr_text for name, dn in wb.defined_names.items()}
    result.set(sheets=info, defined_names=names)
    return result


def cmd_notes(wb, sheet, result):
    out = []
    sheets = [sheet] if sheet else wb.sheetnames
    for name in sheets:
        for row in wb[name].iter_rows():
            for cell in row:
                if cell.comment is not None:
                    out.append({"sheet": name, "cell": cell.coordinate,
                                "text": cell.comment.text,
                                "author": cell.comment.author})
    result.set(notes=out)
    return result


def cmd_names(wb, result):
    names = {name: dn.attr_text for name, dn in wb.defined_names.items()}
    result.set(defined_names=names)
    return result


def cmd_formulas(path, sheet, result):
    """Pass A (formula text) + Pass B (cached values); the two never mix."""
    from openpyxl.utils import coordinate_to_tuple  # noqa: PLC0415

    handle_a = common.load_workbook_safe(path, read_only=True)
    handle_b = common.load_workbook_safe(path, read_only=True, data_only=True)

    cached_map = {}
    for name in handle_b.sheetnames:
        for row_index, row in enumerate(handle_b.iter_values(name), start=1):
            for col_index, value in enumerate(row, start=1):
                if value is not None:
                    cached_map[(name, row_index, col_index)] = value

    sheets = [sheet] if sheet else handle_a.sheetnames
    out = []
    for name in sheets:
        resolve_sheet(handle_a.sheetnames, name, name)
        ws_f = handle_a.wb[name]
        for row in ws_f.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    row_i, col_i = coordinate_to_tuple(cell.coordinate)
                    cached = cached_map.get((name, row_i, col_i))
                    out.append({
                        "sheet": name,
                        "cell": cell.coordinate,
                        "formula": cell.value,
                        "cached": common.jsonable(cached),
                    })
    result.set(formulas=out)
    return result


def main(argv=None):
    ap = argparse.ArgumentParser(description="Read/inspect an .xlsx workbook.")
    ap.add_argument("file", help="path to .xlsx file")
    mode = ap.add_mutually_exclusive_group(required=True)
    mode.add_argument("--sheets", action="store_true")
    mode.add_argument("--json", action="store_true")
    mode.add_argument("--csv", action="store_true")
    mode.add_argument("--formulas", action="store_true")
    mode.add_argument("--notes", action="store_true")
    mode.add_argument("--names", action="store_true")
    ap.add_argument("--sheet", help="sheet name (default: active)")
    ap.add_argument("--data-only", action="store_true",
                    help="return cached formula results (see module docstring)")
    ap.add_argument("--encoding", default="utf-8")
    ap.add_argument("--out", help="output file for --csv")
    args = ap.parse_args(argv)

    result = common.Result(mode="read", file=args.file)

    if args.formulas:
        cmd_formulas(args.file, args.sheet, result)
        return result.emit()

    if args.sheets or args.notes or args.names:
        # Structure inventory: independent of cached values by nature.
        loaded = common.load_workbook_safe(args.file)
        wb = loaded.wb
        if args.sheets:
            cmd_sheets(wb, result)
        elif args.notes:
            if args.sheet:
                resolve_sheet(wb.sheetnames, args.sheet,
                              loaded.active_sheetname)
            cmd_notes(wb, args.sheet, result)
        else:
            cmd_names(wb, result)
        if args.data_only:
            result.diagnose("--data-only has no effect on structure "
                            "inventory (--sheets/--notes/--names).")
        return result.emit()

    loaded = common.load_workbook_safe(args.file, data_only=args.data_only)
    name = resolve_sheet(loaded.sheetnames, args.sheet,
                         loaded.active_sheetname)
    rows = loaded.rows(name)
    if args.json:
        result.set(sheet=name, rows=rows)
        return result.emit()

    # --csv
    if args.out:
        with open(args.out, "w", newline="", encoding=args.encoding) as fh:
            csv.writer(fh).writerows(
                [["" if v is None else v for v in r] for r in rows])
        result.set(out=args.out, rows=len(rows))
        return result.emit()
    writer = csv.writer(sys.stdout)
    for r in rows:
        writer.writerow(["" if v is None else v for v in r])
    return 0


if __name__ == "__main__":
    sys.exit(common.guard(main)())
