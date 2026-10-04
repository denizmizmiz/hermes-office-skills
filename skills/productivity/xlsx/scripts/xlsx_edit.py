#!/usr/bin/env python3
# MIT License. Part of the Hermes xlsx skill.
"""Edit an existing .xlsx workbook in place (or to --out).

Operations (repeatable where noted; execution order is: sheet copies,
sheet renames, structural edits, then cell edits):
  --copy-sheet SRC:NEW          duplicate a sheet under a new name
  --rename-sheet OLD:NEW        rename a sheet
  --insert-rows IDX[:N]         insert N rows before row IDX (default N=1)
  --delete-rows IDX[:N]         delete N rows starting at row IDX
  --insert-cols IDX[:N]         insert N columns before column IDX (number)
  --delete-cols IDX[:N]         delete N columns starting at column IDX
  --set CELL=VALUE              repeatable; type-inferred (int, float, bool,
                                ISO date, else string). '=...' sets a formula.
  --append ROWJSON              repeatable; JSON array appended as a row
  --add-table NAME:RANGE[:STYLE]  create a native Excel table (ListObject)
  --table-append NAME=ROWJSON   append a row inside a table, auto-extending
                                the table's range (repeatable)
  --list-tables                 print tables on the target sheet and exit
  --define-name NAME=REF        workbook-scope defined name, e.g.
                                "Rates='Data'!$B$2:$B$9" (repeatable)
  --delete-name NAME            remove a defined name (repeatable)
  --hyperlink CELL=URL[|TEXT]   set a hyperlink (optional display text)
  --note CELL=TEXT[|AUTHOR]     set a cell note/comment (repeatable)
  --clear-note CELL             remove a cell note (repeatable)
  --protect [PASSWORD]          enable sheet protection; combine with
                                --unlock RANGE to leave ranges editable.
                                NOT security: trivially strippable (see
                                SKILL.md Pitfalls).
  --recalc                      set fullCalcOnLoad so Excel/LibreOffice
                                recomputes all formulas on next open

WARNING: openpyxl does NOT shift merged-cell ranges, chart anchors, or
formula references when rows/columns are inserted or deleted. Verify any
sheet containing merges or formulas after structural edits — or use
xlsx_restructure.py, which rewrites references for you.

Usage:
  xlsx_edit.py book.xlsx --sheet Data --set B2=42 --set C2=2026-01-01 \
      --set "D2==SUM(B2:C2)" --recalc
  xlsx_edit.py book.xlsx --sheet Data --append '["Widget", 9.99, true]'
  xlsx_edit.py book.xlsx --copy-sheet Data:Backup --rename-sheet Data:Main
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

import xlsx_common as common  # noqa: E402

try:  # dependency guard: structured DEPENDENCY_MISSING, never a traceback
    from openpyxl.comments import Comment
    from openpyxl.styles import Protection
    from openpyxl.utils import get_column_letter, range_boundaries
    from openpyxl.workbook.defined_name import DefinedName
    from openpyxl.worksheet.table import Table, TableStyleInfo
except ImportError as _exc:  # pragma: no cover - environment guard
    sys.exit(common.dependency_error(_exc))


def parse_idx(arg):
    if ":" in arg:
        idx, n = arg.split(":", 1)
        return int(idx), int(n)
    return int(arg), 1


def add_table(ws, spec):
    parts = spec.split(":")
    if len(parts) < 3:
        raise ValueError("--add-table needs NAME:RANGE like Sales:A1:C9")
    name = parts[0]
    rng = ":".join(parts[1:3])
    style = parts[3] if len(parts) > 3 else "TableStyleMedium9"
    table = Table(displayName=name, ref=rng)
    table.tableStyleInfo = TableStyleInfo(name=style, showRowStripes=True)
    ws.add_table(table)


def find_table(wb, ws, name, explicit_sheet):
    """Locate a native table by name (D4 fix).

    --sheet given  -> search that one sheet only.
    --sheet absent -> search the active sheet first, then every other sheet.
    Returns (sheet, table) or raises TABLE_NOT_FOUND with the searched list.
    """
    if explicit_sheet:
        if name in ws.tables:
            return ws, ws.tables[name]
        raise common.XlsxError(
            "TABLE_NOT_FOUND",
            f"Table '{name}' not found on sheet '{ws.title}'.",
            recovery="Use --list-tables to see the tables on a sheet.",
            context={"requested_table": name,
                     "searched_sheets": [ws.title],
                     "available_tables": sorted(ws.tables)})
    searched = []
    order = [ws] + [s for s in wb.worksheets if s.title != ws.title]
    for sheet in order:
        searched.append(sheet.title)
        if name in sheet.tables:
            return sheet, sheet.tables[name]
    raise common.XlsxError(
        "TABLE_NOT_FOUND",
        f"Table '{name}' not found on any sheet.",
        recovery="Use --list-tables (with --sheet) to see tables per sheet.",
        context={"requested_table": name, "searched_sheets": searched})


def table_append(ws, name, row_values):
    table = ws.tables[name]
    min_col, min_row, max_col, max_row = range_boundaries(table.ref)
    new_row = max_row + 1
    for offset, value in enumerate(row_values):
        ws.cell(row=new_row, column=min_col + offset, value=value)
    table.ref = (f"{get_column_letter(min_col)}{min_row}:"
                 f"{get_column_letter(max_col)}{new_row}")


def main(argv=None):
    ap = argparse.ArgumentParser(
        description="Edit an existing .xlsx workbook.",
        epilog="Plain insert/delete does not shift merges/formula refs — "
               "use xlsx_restructure.py for reference-aware moves.")
    ap.add_argument("file", help="path to .xlsx file")
    ap.add_argument("--sheet", help="target sheet (default: active)")
    ap.add_argument("--out", help="output path (default: edit in place)")
    ap.add_argument("--rename-sheet", action="append", default=[],
                    metavar="OLD:NEW")
    ap.add_argument("--copy-sheet", action="append", default=[],
                    metavar="SRC:NEW")
    ap.add_argument("--insert-rows", action="append", default=[],
                    metavar="IDX[:N]")
    ap.add_argument("--delete-rows", action="append", default=[],
                    metavar="IDX[:N]")
    ap.add_argument("--insert-cols", action="append", default=[],
                    metavar="IDX[:N]")
    ap.add_argument("--delete-cols", action="append", default=[],
                    metavar="IDX[:N]")
    ap.add_argument("--set", action="append", default=[], metavar="CELL=VALUE")
    ap.add_argument("--append", action="append", default=[], metavar="ROWJSON")
    ap.add_argument("--add-table", action="append", default=[],
                    metavar="NAME:RANGE[:STYLE]")
    ap.add_argument("--table-append", action="append", default=[],
                    metavar="NAME=ROWJSON")
    ap.add_argument("--list-tables", action="store_true",
                    help="print tables on the target sheet and exit")
    ap.add_argument("--define-name", action="append", default=[],
                    metavar="NAME=REF")
    ap.add_argument("--delete-name", action="append", default=[],
                    metavar="NAME")
    ap.add_argument("--hyperlink", action="append", default=[],
                    metavar="CELL=URL[|TEXT]")
    ap.add_argument("--note", action="append", default=[],
                    metavar="CELL=TEXT[|AUTHOR]")
    ap.add_argument("--clear-note", action="append", default=[],
                    metavar="CELL")
    ap.add_argument("--protect", nargs="?", const="", metavar="PASSWORD",
                    help="protect the target sheet (integrity signal only, "
                    "NOT security)")
    ap.add_argument("--unlock", action="append", default=[], metavar="RANGE",
                    help="cell range left editable under --protect")
    ap.add_argument("--recalc", action="store_true",
                    help="force full recalculation when the file is opened")
    ap.add_argument("--approve-token", metavar="TOKEN", default=None,
                    help="approval provenance recorded for in-place writes "
                         "(the approval itself is obtained by the caller)")
    args = ap.parse_args(argv)

    loaded = common.load_workbook_safe(args.file)
    wb = loaded.wb
    changes = []

    # Execution order (D5 fix): copy first, then rename, so a single call can
    # copy a sheet and rename the copy; structural edits come after both.
    for pair in args.copy_sheet:
        src, new = pair.split(":", 1)
        if src not in wb.sheetnames:
            raise common.XlsxError(
                "SHEET_NOT_FOUND",
                f"Worksheet '{src}' does not exist.",
                recovery="Use xlsx_read.py --sheets to list sheet names.",
                context={"requested": src, "available": wb.sheetnames})
        copy = wb.copy_worksheet(wb[src])
        copy.title = new
        changes.append(f"copy {src}->{new}")
    for pair in args.rename_sheet:
        old, new = pair.split(":", 1)
        if old not in wb.sheetnames:
            raise common.XlsxError(
                "SHEET_NOT_FOUND",
                f"Worksheet '{old}' does not exist.",
                recovery="Use xlsx_read.py --sheets to list sheet names.",
                context={"requested": old, "available": wb.sheetnames})
        wb[old].title = new
        changes.append(f"rename {old}->{new}")

    if args.sheet:
        if args.sheet not in wb.sheetnames:
            raise common.XlsxError(
                "SHEET_NOT_FOUND",
                f"Worksheet '{args.sheet}' does not exist.",
                recovery="Use xlsx_read.py --sheets to list sheet names.",
                context={"requested": args.sheet,
                         "available": wb.sheetnames})
        ws = wb[args.sheet]
    else:
        ws = wb.active

    if args.list_tables:
        result = common.Result(mode="edit", sheet=ws.title,
                               tables={t.displayName: {
                                   "ref": t.ref,
                                   "style": t.tableStyleInfo.name
                                   if t.tableStyleInfo else None}
                                   for t in ws.tables.values()})
        return result.emit()

    for arg in args.insert_rows:
        idx, n = parse_idx(arg)
        ws.insert_rows(idx, n)
        changes.append(f"insert_rows {idx}x{n}")
    for arg in args.delete_rows:
        idx, n = parse_idx(arg)
        ws.delete_rows(idx, n)
        changes.append(f"delete_rows {idx}x{n}")
    for arg in args.insert_cols:
        idx, n = parse_idx(arg)
        ws.insert_cols(idx, n)
        changes.append(f"insert_cols {idx}x{n}")
    for arg in args.delete_cols:
        idx, n = parse_idx(arg)
        ws.delete_cols(idx, n)
        changes.append(f"delete_cols {idx}x{n}")

    for assignment in args.set:
        coord, raw = assignment.split("=", 1)
        ws[coord] = common.infer(raw, keep_empty_string=True)
        changes.append(f"set {coord}")
    for row_json in args.append:
        ws.append(json.loads(row_json))
        changes.append(f"append row {ws.max_row}")

    for spec in args.add_table:
        add_table(ws, spec)
        changes.append(f"add_table {spec.split(':')[0]}")
    for spec in args.table_append:
        name, row_json = spec.split("=", 1)
        owner, _table = find_table(wb, ws, name, bool(args.sheet))
        table_append(owner, name, json.loads(row_json))
        changes.append(f"table_append {name} -> {owner.tables[name].ref}")

    for spec in args.define_name:
        name, ref = spec.split("=", 1)
        wb.defined_names[name] = DefinedName(name, attr_text=ref)
        changes.append(f"define_name {name}")
    for name in args.delete_name:
        del wb.defined_names[name]
        changes.append(f"delete_name {name}")

    for spec in args.hyperlink:
        coord, rest = spec.split("=", 1)
        url, _, text = rest.partition("|")
        cell = ws[coord]
        cell.hyperlink = url
        cell.value = text or (cell.value if cell.value is not None else url)
        cell.style = "Hyperlink"
        changes.append(f"hyperlink {coord}")
    for spec in args.note:
        coord, rest = spec.split("=", 1)
        text, _, author = rest.partition("|")
        ws[coord].comment = Comment(text, author or "xlsx-skill")
        changes.append(f"note {coord}")
    for coord in args.clear_note:
        ws[coord].comment = None
        changes.append(f"clear_note {coord}")

    if args.protect is not None:
        for rng in args.unlock:
            for row in ws[rng]:
                for cell in row:
                    cell.protection = Protection(locked=False)
        if args.protect:
            ws.protection.password = args.protect
        ws.protection.sheet = True
        changes.append(f"protect {ws.title}"
                       + (f" (unlocked {len(args.unlock)} ranges)"
                          if args.unlock else ""))

    if args.recalc:
        wb.calculation.fullCalcOnLoad = True
        changes.append("fullCalcOnLoad")

    # Safe write (F): backup -> temp -> validate -> atomic commit, all via
    # xlsx_common.save_workbook_safe. Legacy in-place default (--out omitted)
    # is kept and recorded as caller-explicit; --out writes a new file.
    target = args.out or args.file
    in_place = (Path(target).resolve() == Path(args.file).resolve())
    save_info = common.save_workbook_safe(
        loaded, target, backup=True, in_place=in_place,
        approve_token=args.approve_token,
        expect_sheets=list(wb.sheetnames),
        plan={"action": "edit", "sheet": ws.title, "changes": changes})

    result = common.Result(mode="edit", sheet=ws.title, changes=changes)
    result.set(output=save_info["output"], backup=save_info["backup"],
               atomic=save_info["atomic"], bytes=save_info["bytes"],
               operation_id=save_info["operation_id"],
               plan_hash=save_info["plan_hash"],
               approval=save_info["approval"],
               manifest=save_info["manifest"])
    if save_info["manifest"] and not save_info["manifest"].get("ok"):
        result.warn("FEATURE_PARTIAL",
                    "the write manifest could not be updated: "
                    + str(save_info["manifest"].get("message")))
    requested = any([args.copy_sheet, args.rename_sheet, args.insert_rows,
                     args.delete_rows, args.insert_cols, args.delete_cols,
                     args.set, args.append, args.add_table, args.table_append,
                     args.define_name, args.delete_name, args.hyperlink,
                     args.note, args.clear_note, args.protect is not None,
                     args.recalc])
    if not requested:
        result.warn("SILENT_NOOP_PREVENTED",
                    "No edit operations were requested; the workbook was "
                    "rewritten unchanged.")
    if save_info["backup"]:
        result.diagnose(
            f"previous file backed up to {save_info['backup']}")
    return result.emit()


if __name__ == "__main__":
    sys.exit(common.guard(main)())
