#!/usr/bin/env python3
# MIT License. Part of the Hermes xlsx skill.
"""Convert a CSV file to a styled .xlsx workbook with type inference.

Type inference per cell (disable with --no-infer):
  int, float, bool ("true"/"false", case-insensitive), ISO date
  (YYYY-MM-DD) and ISO datetime; everything else stays a string.

Styling applied by default (disable with --plain):
  bold header row with a light fill, frozen top row, autofilter over the
  data range, and column widths sized to the longest cell (capped at 60).

The output goes through the shared safe-write chain (backup -> temp ->
validate -> atomic commit); an existing file at the output path is backed
up first.

Usage:
  csv_to_xlsx.py data.csv out.xlsx
  csv_to_xlsx.py data.csv out.xlsx --sheet-name Import --encoding cp1252
  csv_to_xlsx.py data.csv out.xlsx --delimiter ';' --no-infer
"""
from __future__ import annotations

import argparse
import csv
import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

import xlsx_common as common  # noqa: E402

try:  # dependency guard: structured DEPENDENCY_MISSING, never a traceback
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError as _exc:  # pragma: no cover - environment guard
    sys.exit(common.dependency_error(_exc))

MAX_COL_WIDTH = 60
COL_PADDING = 2
DEFAULT_COL_WIDTH = 8


def main(argv=None):
    ap = argparse.ArgumentParser(description="CSV -> styled .xlsx converter.")
    ap.add_argument("csv_file", help="input CSV path")
    ap.add_argument("output", help="output .xlsx path")
    ap.add_argument("--sheet-name", default="Sheet1")
    ap.add_argument("--encoding", default="utf-8",
                    help="CSV file encoding (default utf-8)")
    ap.add_argument("--delimiter", default=",")
    ap.add_argument("--no-infer", action="store_true",
                    help="keep every cell as a string")
    ap.add_argument("--plain", action="store_true",
                    help="skip header styling / freeze / autofilter")
    args = ap.parse_args(argv)

    csv_path = Path(args.csv_file)
    if not csv_path.exists():
        raise common.XlsxError(
            "FILE_NOT_FOUND",
            f"No such file: {args.csv_file}",
            recovery="Check the path and try again.",
            context={"path": str(csv_path)})

    try:
        with open(args.csv_file, newline="", encoding=args.encoding) as fh:
            rows = list(csv.reader(fh, delimiter=args.delimiter))
    except UnicodeDecodeError as exc:
        raise common.XlsxError(
            "UNEXPECTED_ERROR",
            f"CSV could not be decoded as {args.encoding}: {exc}",
            recovery="Pass the correct --encoding (e.g. cp1252, utf-8-sig).",
            context={"path": str(csv_path), "encoding": args.encoding}) from exc

    wb = Workbook()
    ws = wb.active
    ws.title = args.sheet_name
    for i, row in enumerate(rows):
        if args.no_infer or i == 0:
            ws.append(row)
        else:
            ws.append([common.infer(cell) for cell in row])

    if rows and not args.plain:
        header_font = Font(bold=True)
        header_fill = PatternFill("solid", fgColor="DDEBF7")
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for col_idx in range(1, ws.max_column + 1):
            longest = max((len(str(r[col_idx - 1])) for r in rows
                           if len(r) >= col_idx), default=DEFAULT_COL_WIDTH)
            ws.column_dimensions[get_column_letter(col_idx)].width = \
                min(longest + COL_PADDING, MAX_COL_WIDTH)

    # Safe write (F): fresh workbook through the shared chain.
    loaded = common.LoadedWorkbook(wb, path=None, mode="normal")
    save_info = common.save_workbook_safe(
        loaded, args.output, backup=True, expect_sheets=wb.sheetnames,
        plan={"action": "csv_to_xlsx", "sheet": args.sheet_name,
              "rows": len(rows)})

    result = common.Result(mode="csv_to_xlsx", rows=len(rows))
    result.set(output=save_info["output"], backup=save_info["backup"],
               atomic=save_info["atomic"], bytes=save_info["bytes"],
               operation_id=save_info["operation_id"],
               plan_hash=save_info["plan_hash"],
               approval=save_info["approval"],
               manifest=save_info["manifest"])
    if save_info["manifest"] and not save_info["manifest"].get("ok"):
        result.warn("FEATURE_PARTIAL",
                    "the write manifest could not be updated: "
                    + str(save_info["manifest"].get("message")))
    if save_info["backup"]:
        result.diagnose(
            f"an existing file at the output path was backed up to "
            f"{save_info['backup']}")
    return result.emit()


if __name__ == "__main__":
    sys.exit(common.guard(main)())
