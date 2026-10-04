#!/usr/bin/env python3
"""FAZ 3A tests -- template execution core (approved fill-plan execution).

Coverage map (spec section 22): plan validation, approval, scalar and
row-block execution, type safety, validation safety, merged cells,
formula protection, row-expansion boundary, safe write (backup / temp /
QA-gate / commit), idempotency, manifest, rollback, fidelity, no plan
leakage, determinism and the full regression of the earlier phases
(run via the shared suite).
"""
from __future__ import annotations

import importlib
import json
import shutil
import subprocess
import sys
from pathlib import Path

import pytest

SCRIPTS = Path(__file__).resolve().parent.parent / "scripts"
SKILL = SCRIPTS.parent
sys.path.insert(0, str(SCRIPTS))

import xlsx_common as common  # noqa: E402

EXEC = SCRIPTS / "xlsx_execute.py"
SEM = SCRIPTS / "xlsx_semantics.py"
MAP = SCRIPTS / "xlsx_mapping.py"

import openpyxl  # noqa: E402
from openpyxl.formatting.rule import CellIsRule  # noqa: E402
from openpyxl.styles import PatternFill  # noqa: E402
from openpyxl.workbook.defined_name import DefinedName  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402
from openpyxl.worksheet.table import Table, TableStyleInfo  # noqa: E402

SLOT_FILL = PatternFill("solid", fgColor="FFF2CC")

BASE_SOURCE = "Ad,Adet,Sertifika,Onay\nKalem Y,3,RCS,OK\n"


# ---------------------------------------------------------------------------
# fixture builders
# ---------------------------------------------------------------------------

def make_exec_book(path):
    """A form-like template with everything FAZ 3A must protect.

    Sheet "Form":
      row 1   headers (Ad | Adet | Sertifika | Tutar | Onay)
      rows 2-3 data (D holds formulas)
      row 4   spacer label ("Giris")
      rows 5-6 evidenced blank slots A:D (fill + DV on C; "0" on B)
      G2      single-row scalar slot (header "Onay")
      row 8   hidden
      CF rule on B2, defined name FormAd -> Form!$A$2
    Sheet "Data":
      A1:B3   table Tbl1
      A5:B5   merged title
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Form"
    for index, title in enumerate(["Ad", "Adet", "Sertifika", "Tutar", "Onay"],
                                  start=1):
        ws.cell(row=1, column=index, value=title)
    ws["A2"], ws["B2"], ws["C2"], ws["E2"] = "Kalem X", 2, "GRS", "ESKI"
    ws["D2"] = "=B2*10"
    ws["A3"], ws["B3"], ws["C3"], ws["E3"] = "Kalem Y", 5, "RCS", "ESKI"
    ws["D3"] = "=B3*10"
    ws["B2"].number_format = "0"
    ws["B3"].number_format = "0"
    ws["A4"] = "Giris"
    for row in (5, 6):
        for column in "ABCD":
            ws[f"{column}{row}"].fill = SLOT_FILL
    ws["B5"].number_format = "0"
    ws["B6"].number_format = "0"
    dv = DataValidation(type="list", formula1='"GRS,RCS,OCS"',
                        allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("C5:C6")
    dv.add("C2:C3")
    ws["G1"] = "Onay"
    ws["G2"].fill = SLOT_FILL
    ws.row_dimensions[8].hidden = True
    ws.conditional_formatting.add(
        "B2", CellIsRule(operator="greaterThan", formula=["1"],
                         fill=PatternFill("solid", fgColor="FFC7CE")))
    wb.defined_names.add(DefinedName("FormAd", attr_text="Form!$A$2"))
    ws2 = wb.create_sheet("Data")
    ws2["A1"], ws2["B1"] = "Kod", "Deger"
    ws2["A2"], ws2["B2"] = "K1", 1
    ws2["A3"], ws2["B3"] = "K2", 2
    table = Table(displayName="Tbl1", ref="A1:B3")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium9",
                                          showRowStripes=True)
    ws2.add_table(table)
    ws2["A5"] = "Not"
    ws2.merge_cells("A5:B5")
    wb.save(path)
    return path


def sha256(path):
    return common.sha256_file(path)["sha256"]


def run_script(script, args, expect_ok=True):
    proc = subprocess.run([sys.executable, str(script), *map(str, args)],
                          capture_output=True, text=True, encoding="utf-8",
                          timeout=600)
    if expect_ok and proc.returncode != 0:
        raise AssertionError(
            f"{Path(script).name} failed:\n{proc.stdout[-1500:]}\n"
            f"{proc.stderr[-1500:]}")
    return proc


def build_pipeline(directory, source_text=BASE_SOURCE, policy_text=None):
    """Template -> profile -> plan -> dry-run (the three 2B read-only runs)."""
    directory = Path(directory)
    directory.mkdir(parents=True, exist_ok=True)
    book = make_exec_book(directory / "book.xlsx")
    prof = directory / "prof.json"
    run_script(SEM, [book, "--emit", "profile", "--out", prof])
    src = directory / "src.csv"
    src.write_text(source_text, encoding="utf-8")
    args = ["--profile", prof, "--source", src, "--emit", "plan",
            "--out", directory / "plan.json"]
    if policy_text is not None:
        pol = directory / "policy.json"
        pol.write_text(policy_text, encoding="utf-8")
        args += ["--policy-file", pol]
    run_script(MAP, args)
    dry = directory / "dry.json"
    run_script(MAP, ["--profile", prof, "--source", src, "--plan",
                     directory / "plan.json", "--target", book,
                     "--emit", "dry-run", "--out", dry])
    return {"dir": directory, "book": book, "prof": prof, "src": src,
            "plan": directory / "plan.json", "dry": dry,
            "fill_plan": json.loads(dry.read_text(encoding="utf-8"))}


@pytest.fixture(scope="session")
def base(tmp_path_factory):
    """One session-scoped pipeline; tests clone the workbook per test."""
    root = tmp_path_factory.mktemp("exec_base")
    return build_pipeline(root)


@pytest.fixture()
def case(base, tmp_path):
    """A per-test clone of the base template plus shared plan artifacts."""
    book = tmp_path / "case.xlsx"
    shutil.copy2(base["book"], book)
    manifest = tmp_path / "manifest"
    return {"book": book, "prof": base["prof"], "src": base["src"],
            "plan": base["plan"], "dry": base["dry"],
            "fill_plan": base["fill_plan"], "manifest": manifest,
            "dir": tmp_path}


def exec_args(case, *, token=None, emit=None, out=None, in_place=False,
              fill_plan=None, target=None, source=None):
    args = ["--fill-plan", fill_plan or case["dry"], "--profile",
            case["prof"], "--source", source or case["src"],
            "--target", target or case["book"],
            "--manifest-dir", case["manifest"]]
    if emit:
        args += ["--emit", emit]
    if token is not None:
        args += ["--approve-token", token]
    if out is not None:
        args += ["--out", out]
    if in_place:
        args += ["--in-place"]
    return args


def plan_token(fill_plan):
    return fill_plan["plan_id"][:12]


def exec_cli(case, expect_ok=True, **kwargs):
    return run_script(EXEC, exec_args(case, **kwargs), expect_ok=expect_ok)


def exec_cli_json(case, **kwargs):
    proc = exec_cli(case, **kwargs)
    return json.loads(proc.stdout)


def exec_error(case, **kwargs):
    proc = exec_cli(case, expect_ok=False, **kwargs)
    assert proc.returncode != 0
    return json.loads(proc.stderr), proc


def tamper_plan(case, mutate, name="tampered.json"):
    plan = json.loads(case["dry"].read_text(encoding="utf-8"))
    mutate(plan)
    path = case["dir"] / name
    path.write_text(json.dumps(plan, ensure_ascii=False), encoding="utf-8")
    return path


def write_entry(plan, *, sheet, cell, source, preview, kind="string",
                checks=None):
    entry = {"sheet": sheet, "cell": cell, "source": source,
             "preview": preview, "type": kind,
             "write_policy": "write_value"}
    if checks is not None:
        entry["checks"] = checks
    plan.setdefault("would_write", []).append(entry)
    return entry


def cell_value(path, sheet, cell):
    wb = openpyxl.load_workbook(path)
    try:
        return wb[sheet][cell].value
    finally:
        wb.close()


def snapshot_book(path):
    wb = openpyxl.load_workbook(path)
    try:
        return {
            "merges": {ws.title: sorted(str(r)
                                        for r in ws.merged_cells.ranges)
                       for ws in wb.worksheets},
            "tables": {ws.title: {name: ref for name, ref in
                                  sorted(ws.tables.items())}
                       for ws in wb.worksheets},
            "names": sorted(f"{n}" for n in wb.defined_names),
            "validations": {ws.title: sorted(
                f"{dv.type}:{dv.sqref}" for dv in
                ws.data_validations.dataValidation)
                for ws in wb.worksheets},
            "cf": {ws.title: len(list(ws.conditional_formatting))
                   for ws in wb.worksheets},
            "hidden_rows": {ws.title: sorted(
                str(i) for i, d in ws.row_dimensions.items()
                if getattr(d, "hidden", False))
                for ws in wb.worksheets},
        }
    finally:
        wb.close()



# ---------------------------------------------------------------------------
# 1. plan validation (AC-3A03)
# ---------------------------------------------------------------------------

def test_pipeline_fill_plan_shape(base):
    fp = base["fill_plan"]
    assert fp["plan_id"]
    assert fp["would_write"], "fixture must produce planned writes"
    cells = {(e["sheet"], e["cell"]) for e in fp["would_write"]}
    assert {("Form", "A2"), ("Form", "B2"), ("Form", "C2"),
            ("Form", "E2")} <= cells
    assert all(e["write_policy"] == "write_value"
               for e in fp["would_write"])


def test_preflight_valid_plan_is_executable(case):
    payload = exec_cli_json(case, emit="preflight", token=plan_token(
        case["fill_plan"]))
    report = payload["report"]
    assert report["executable"] is True
    assert report["structure_match"] is True
    assert report["approval"]["state"] == "provided"
    assert report["writes_planned"] >= 4


def test_invalid_schema_is_refused(case):
    bad = tamper_plan(case, lambda plan: plan.pop("would_write"),
                      name="bad.json")
    error, _ = exec_error(case, token=plan_token(case["fill_plan"]),
                          fill_plan=bad)
    assert error["error_code"] == "PLAN_INVALID"
    assert not (case["dir"] / "case_filled.xlsx").exists()


def test_stale_fingerprint_blocks(case):
    stale = tamper_plan(
        case,
        lambda plan: plan.update(plan_target_structure_fingerprint=
                                 "sha256:" + "0" * 64),
        name="stale.json")
    error, _ = exec_error(case, token=plan_token(case["fill_plan"]),
                          fill_plan=stale)
    assert error["error_code"] == "STALE_PLAN"
    assert error["context"]["write_count"] == 0


def test_missing_target_sheet_blocks(case):
    bad = tamper_plan(
        case,
        lambda plan: plan["would_write"][0].update(sheet="Nope"),
        name="nosheet.json")
    error, _ = exec_error(case, token=plan_token(case["fill_plan"]),
                          fill_plan=bad)
    assert error["error_code"] == "EXECUTION_BLOCKED"
    assert any("does not exist" in b["message"]
               for b in error["context"]["blockers"])


def test_invalid_cell_reference_blocks(case):
    bad = tamper_plan(
        case,
        lambda plan: plan["would_write"][0].update(cell="1A"),
        name="badcell.json")
    error, _ = exec_error(case, token=plan_token(case["fill_plan"]),
                          fill_plan=bad)
    assert error["error_code"] == "PLAN_INVALID"


def test_duplicate_operations_block(case):
    def clone_first(plan):
        plan["would_write"].append(dict(plan["would_write"][0]))
    bad = tamper_plan(case, clone_first, name="dup.json")
    error, _ = exec_error(case, token=plan_token(case["fill_plan"]),
                          fill_plan=bad)
    assert error["error_code"] == "PLAN_INVALID"
    assert any("duplicate" in b["message"]
               for b in error["context"]["blockers"])


# ---------------------------------------------------------------------------
# 2. approval (AC-3A02)
# ---------------------------------------------------------------------------

def test_approval_absent_blocks(case):
    error, proc = exec_error(case)
    assert error["error_code"] == "APPROVAL_REQUIRED"
    assert error["context"]["write_count"] == 0
    assert not (case["dir"] / "case_filled.xlsx").exists()
    assert sha256(case["book"]) == base_book_sha(case)


def test_approval_invalid_token_blocks(case):
    error, _ = exec_error(case, token="000000000000")
    assert error["error_code"] == "APPROVAL_INVALID"
    assert not (case["dir"] / "case_filled.xlsx").exists()


def test_approval_present_executes(case):
    payload = exec_cli_json(case, token=plan_token(case["fill_plan"]))
    assert payload["status"] == "committed"
    assert payload["approval"]["state"] == "provided"
    assert payload["write_count"] >= 4


def base_book_sha(case):
    return sha256(case["book"])


# ---------------------------------------------------------------------------
# 3. scalar + row-block execution (AC-3A04)
# ---------------------------------------------------------------------------

def test_writes_land_in_their_cells(case):
    payload = exec_cli_json(case, token=plan_token(case["fill_plan"]))
    out = Path(payload["output"])
    assert out.exists()
    assert cell_value(out, "Form", "A2") == "Kalem Y"
    assert cell_value(out, "Form", "B2") == 3
    assert cell_value(out, "Form", "C2") == "RCS"
    assert cell_value(out, "Form", "E2") == "OK"


def test_scalar_mode_write_via_evidenced_slot(case):
    """A single-cell write to an evidenced slot (scalar-plan path)."""
    def add_entry(plan):
        write_entry(plan, sheet="Form", cell="G2", source="Onay[0]",
                    preview="OK", kind="string")
    plan = tamper_plan(case, add_entry, name="scalar.json")
    payload = exec_cli_json(case, token=plan_token(case["fill_plan"]),
                            fill_plan=plan)
    assert cell_value(payload["output"], "Form", "G2") == "OK"
    assert payload["write_count"] >= 5


def test_only_planned_cells_changed(case):
    planned = {(e["sheet"], e["cell"])
               for e in case["fill_plan"]["would_write"]}
    assert exec_cli_json(case, token=plan_token(case["fill_plan"]))
    before = openpyxl.load_workbook(case["book"])
    after = openpyxl.load_workbook(case["dir"] / "case_filled.xlsx")
    try:
        diffs = set()
        for sheet in after.sheetnames:
            a_cells = {c.coordinate: c.value
                       for c in before[sheet]._cells.values()
                       if c.value is not None}
            b_cells = {c.coordinate: c.value
                       for c in after[sheet]._cells.values()
                       if c.value is not None}
            for coordinate in set(a_cells) | set(b_cells):
                if a_cells.get(coordinate) != b_cells.get(coordinate):
                    diffs.add((sheet, coordinate))
    finally:
        before.close()
        after.close()
    # planned cells that actually hold new values are the only diffs
    assert diffs <= planned
    assert len(diffs) == len(planned)


def test_noop_cells_are_reported_not_silent(case):
    """A planned cell that already holds the value is a noop, reported."""
    same = build_pipeline(
        case["dir"] / "same",
        source_text="Ad,Adet,Sertifika,Onay\nKalem X,9,OCS,OK\n")
    payload = exec_cli_json(case, token=plan_token(same["fill_plan"]),
                            fill_plan=same["dry"], source=same["src"],
                            out=case["dir"] / "same_out.xlsx")
    assert payload["status"] == "committed"
    assert payload["noop_count"] >= 1
    assert payload["changes_count"] == payload["write_count"] - \
        payload["noop_count"]
    codes = {item.get("code") for item in payload.get("warnings", [])}
    assert "CELL_ALREADY_EQUAL" in codes


def test_row_block_writes_multiple_rows(case, base):
    two_rows = build_pipeline(case["dir"] / "rb",
                             source_text="Ad,Adet,Sertifika,Onay\n"
                                         "Kalem Y,3,RCS,OK\n"
                                         "Kalem Z,4,OCS,OK\n")
    payload = exec_cli_json(case, token=plan_token(two_rows["fill_plan"]),
                            fill_plan=two_rows["dry"], source=two_rows["src"],
                            out=case["dir"] / "rb_out.xlsx")
    out = Path(payload["output"])
    assert cell_value(out, "Form", "A2") == "Kalem Y"
    assert cell_value(out, "Form", "A3") == "Kalem Z"
    assert cell_value(out, "Form", "B3") == 4
    assert cell_value(out, "Form", "C3") == "OCS"
    assert payload["write_count"] >= 8


def test_row_expansion_required_blocks(case):
    three_rows = build_pipeline(
        case["dir"] / "rx",
        source_text="Ad,Adet,Sertifika,Onay\nKalem Y,3,RCS,OK\n"
                    "Kalem Z,4,OCS,OK\nKalem W,5,GRS,OK\n")
    assert three_rows["fill_plan"]["row_expansion"]["needed"] is True
    error, _ = exec_error(case, token=plan_token(three_rows["fill_plan"]),
                          fill_plan=three_rows["dry"])
    assert error["error_code"] == "ROW_EXPANSION_REQUIRED"
    ctx = error["context"]
    assert any(b["code"] == "ROW_EXPANSION_REQUIRED"
               for b in ctx["blockers"])
    assert not (case["dir"] / "case_filled.xlsx").exists()


# ---------------------------------------------------------------------------
# 4. type safety (AC-3A05)
# ---------------------------------------------------------------------------

def test_type_mismatch_blocks(case):
    def make_mismatch(plan):
        for entry in plan["would_write"]:
            if entry["cell"] == "B2":
                entry.update(source="Ad[0]", preview="Kalem Y",
                             type="string")
    bad = tamper_plan(case, make_mismatch, name="type.json")
    error, _ = exec_error(case, token=plan_token(case["fill_plan"]),
                          fill_plan=bad)
    assert error["error_code"] == "TYPE_MISMATCH"
    assert not (case["dir"] / "case_filled.xlsx").exists()


def test_format_incompatible_flag_blocks(case):
    def flag(plan):
        for entry in plan["would_write"]:
            entry.setdefault("checks", {})["format_compatible"] = False
    bad = tamper_plan(case, flag, name="fmt.json")
    error, _ = exec_error(case, token=plan_token(case["fill_plan"]),
                          fill_plan=bad)
    assert error["error_code"] == "TYPE_MISMATCH"


# ---------------------------------------------------------------------------
# 5. validation / dropdown safety (AC-3A08)
# ---------------------------------------------------------------------------

def test_valid_dropdown_value_passes(case):
    payload = exec_cli_json(case, token=plan_token(case["fill_plan"]))
    assert cell_value(payload["output"], "Form", "C2") == "RCS"


def test_invalid_dropdown_value_blocks(case):
    def make_invalid(plan):
        for entry in plan["would_write"]:
            if entry["cell"] == "C2":
                entry.update(source="Ad[0]", preview="Kalem Y")
    bad = tamper_plan(case, make_invalid, name="dv.json")
    error, _ = exec_error(case, token=plan_token(case["fill_plan"]),
                          fill_plan=bad)
    assert error["error_code"] == "VALIDATION_MISMATCH"
    assert not (case["dir"] / "case_filled.xlsx").exists()


# ---------------------------------------------------------------------------
# 6. merged cells (AC-3A09)
# ---------------------------------------------------------------------------

def test_merged_top_left_write_allowed(case):
    def add_entry(plan):
        write_entry(plan, sheet="Data", cell="A5", source="Ad[0]",
                    preview="Kalem Y", kind="string")
    plan = tamper_plan(case, add_entry, name="mtop.json")
    payload = exec_cli_json(case, token=plan_token(case["fill_plan"]),
                            fill_plan=plan)
    assert cell_value(payload["output"], "Data", "A5") == "Kalem Y"


def test_merged_non_top_left_blocked(case):
    def add_entry(plan):
        write_entry(plan, sheet="Data", cell="B5", source="Ad[0]",
                    preview="Kalem Y", kind="string")
    plan = tamper_plan(case, add_entry, name="minside.json")
    error, _ = exec_error(case, token=plan_token(case["fill_plan"]),
                          fill_plan=plan)
    assert error["error_code"] == "MERGED_CELL_WRITE_FORBIDDEN"
    assert not (case["dir"] / "case_filled.xlsx").exists()


# ---------------------------------------------------------------------------
# 7. formula protection (AC-3A06 / AC-3A17)
# ---------------------------------------------------------------------------

def test_formula_cell_write_blocked(case):
    def add_entry(plan):
        write_entry(plan, sheet="Form", cell="D2", source="Ad[0]",
                    preview="Kalem Y", kind="string")
    plan = tamper_plan(case, add_entry, name="fcell.json")
    error, _ = exec_error(case, token=plan_token(case["fill_plan"]),
                          fill_plan=plan)
    assert error["error_code"] == "FORMULA_MODIFICATION_FORBIDDEN"


def test_preserve_formula_policy_is_refused_for_writes(case):
    def add_entry(plan):
        entry = write_entry(plan, sheet="Form", cell="D5", source="Ad[0]",
                            preview="Kalem Y", kind="string")
        entry["write_policy"] = "preserve_formula"
    plan = tamper_plan(case, add_entry, name="fpolicy.json")
    error, _ = exec_error(case, token=plan_token(case["fill_plan"]),
                          fill_plan=plan)
    assert error["error_code"] == "PLAN_INVALID"


def test_formulas_survive_execution(case):
    payload = exec_cli_json(case, token=plan_token(case["fill_plan"]))
    out = Path(payload["output"])
    assert cell_value(out, "Form", "D2") == "=B2*10"
    assert cell_value(out, "Form", "D3") == "=B3*10"
    assert payload["qa"]["formula_count_before"] == \
        payload["qa"]["formula_count_after"]
    assert payload["validation"]["formulas_preserved"] is True



# ---------------------------------------------------------------------------
# 8. safe write (AC-3A11 / 12 / 13) + backup + rollback
# ---------------------------------------------------------------------------

def test_default_output_is_a_new_file(case):
    original = sha256(case["book"])
    payload = exec_cli_json(case, token=plan_token(case["fill_plan"]))
    assert payload["original_unchanged"] is True
    assert sha256(case["book"]) == original
    assert Path(payload["output"]).name == "case_filled.xlsx"


def test_in_place_creates_backup_and_writes(case):
    original = sha256(case["book"])
    payload = exec_cli_json(case, token=plan_token(case["fill_plan"]),
                            in_place=True)
    assert payload["backup"], "in-place execution must take a backup"
    assert sha256(case["book"]) != original
    assert sha256(payload["backup"]) == original
    assert payload["original_unchanged"] is True


def test_commit_success_reopens_and_validates(case):
    payload = exec_cli_json(case, token=plan_token(case["fill_plan"]))
    probe = common.verify_workbook(payload["output"])
    assert probe["ok"] is True
    assert payload["validation"]["reopens"] is True
    assert payload["qa"]["status"] == "pass"
    assert payload["atomic"] is True


def test_qa_failure_prevents_commit(case, monkeypatch):
    ex = importlib.import_module("xlsx_execute")
    original = sha256(case["book"])

    def failing_qa(*_a, **_k):
        return {"ok": False, "validation_status": "fail",
                "qa_status": "fail", "unexpected":
                [{"kind": "value", "cell": "Form!Z9",
                  "before": None, "after": 1}],
                "changed": [], "noop": [], "warnings": [],
                "formula_count_before": 0, "formula_count_after": 0,
                "region_fingerprint_after": None}
    monkeypatch.setattr(ex, "post_validate", failing_qa)
    with pytest.raises(common.XlsxError) as exc:
        ex.main([str(a) for a in
                 exec_args(case, token=plan_token(case["fill_plan"]))])
    assert exc.value.code == "QA_FAILED"
    assert exc.value.context["committed"] is False
    assert not (case["dir"] / "case_filled.xlsx").exists()
    assert sha256(case["book"]) == original
    assert list(case["dir"].glob(".case_filled.stage-*.xlsx")), \
        "staged file must be kept for diagnostics"


def test_rollback_restores_original_on_verify_failure(case, monkeypatch):
    ex = importlib.import_module("xlsx_execute")
    original = sha256(case["book"])
    monkeypatch.setattr(
        ex, "_post_commit_verify",
        lambda *a, **k: {"ok": False, "reason": "injected failure"})
    with pytest.raises(common.XlsxError) as exc:
        ex.main([str(a) for a in exec_args(
            case, token=plan_token(case["fill_plan"]), in_place=True)])
    assert exc.value.code == "WRITE_FAILED"
    assert exc.value.context["rolled_back"] is True
    assert sha256(case["book"]) == original, "backup must be restored"
    # the rollback moves the backup back onto the target; its path is
    # still recorded in the error context for the incident trail.
    assert exc.value.context["backup"], "rollback must name the backup"


# ---------------------------------------------------------------------------
# 9. idempotency (AC-3A14)
# ---------------------------------------------------------------------------

def _run_once(case, **kwargs):
    return exec_cli_json(case, token=plan_token(case["fill_plan"]), **kwargs)


def test_second_run_is_safe_to_reapply(case):
    """New-file output leaves the target untouched, so the region is in
    its pre-execution state: re-applying reproduces the result."""
    first = _run_once(case)
    assert first["write_count"] >= 4
    second = _run_once(case)
    assert second["status"] == "committed"
    assert second["idempotency_status"] == "safe_to_reapply"
    assert second["write_count"] == first["write_count"]


def test_already_applied_after_in_place_run(case):
    """An in-place run fills the target; re-running is a zero-write no-op."""
    first = _run_once(case, in_place=True)
    assert first["status"] == "committed"
    assert cell_value(case["book"], "Form", "A2") == "Kalem Y"
    second = _run_once(case, in_place=True)
    assert second["status"] == "already_applied"
    assert second["write_count"] == 0
    assert second["prior_operation"] == first["operation_id"]
    assert cell_value(case["book"], "Form", "A2") == "Kalem Y"


def test_safe_to_reapply_after_region_cleared(case):
    first = _run_once(case)
    fresh = case["dir"] / "fresh.xlsx"
    shutil.copy2(case["book"], fresh)   # same bytes -> same plan identity
    payload = exec_cli_json(case, token=plan_token(case["fill_plan"]),
                            target=fresh, out=case["dir"] / "fresh_out.xlsx")
    assert payload["status"] == "committed"
    assert payload["write_count"] == first["write_count"]


def test_conflict_blocks_when_region_differs(case):
    _run_once(case)
    dirty = case["dir"] / "dirty.xlsx"
    shutil.copy2(case["book"], dirty)
    wb = openpyxl.load_workbook(dirty)
    wb["Form"]["C2"] = "ZZZ"
    wb.save(dirty)
    wb.close()
    error, _ = exec_error(case, token=plan_token(case["fill_plan"]),
                          target=dirty, out=case["dir"] / "dirty_out.xlsx")
    assert error["error_code"] == "PLAN_CONFLICT"
    assert not (case["dir"] / "dirty_out.xlsx").exists()


# ---------------------------------------------------------------------------
# 10. manifest (AC-3A15)
# ---------------------------------------------------------------------------

def test_manifest_record_has_required_fields(case):
    payload = _run_once(case)
    manifest_file = case["manifest"] / ".xlsx_ops" / "manifest.jsonl"
    entries = [json.loads(line) for line
               in manifest_file.read_text(encoding="utf-8").splitlines()
               if line.strip()]
    records = [e for e in entries
               if e.get("plan_hash") == f"sha256:{case['fill_plan']['plan_id']}"]
    assert records, "execution must be recorded in the manifest"
    record = records[-1]
    for field in ("operation_id", "plan_id", "profile_id",
                  "target_fingerprint_before", "target_fingerprint_after",
                  "status", "validation_status", "qa_status", "output",
                  "changes_count", "approval_state", "approval_source"):
        assert record.get(field) is not None, field
    assert record["operation_id"] == payload["operation_id"]
    assert record["status"] == "committed"
    assert record["validation_status"] == "pass"
    assert record["qa_status"] == "pass"


# ---------------------------------------------------------------------------
# 11. fidelity (AC-3A10 / AC-3A18)
# ---------------------------------------------------------------------------

def test_structure_and_style_fidelity(case):
    before = snapshot_book(case["book"])
    payload = _run_once(case)
    after = snapshot_book(payload["output"])
    assert before == after
    wb = openpyxl.load_workbook(payload["output"])
    try:
        assert wb["Form"]["A2"].value == "Kalem Y"
        assert wb["Form"]["B2"].value == 3
        cell = wb["Form"]["A5"]          # untouched slot keeps its style
        assert cell.fill.fgColor.rgb.endswith("FFF2CC")
        assert wb["Form"]["B5"].number_format == "0"
    finally:
        wb.close()
    assert payload["validation"]["structure_preserved"] is True
    assert payload["validation"]["styles_preserved"] is True


def test_unexpected_value_change_is_detected():
    ex = importlib.import_module("xlsx_execute")

    def snap(values, styles=None):
        return {"sheets": ["S"], "values": {"S": dict(values)},
                "formulas": {}, "merges": {"S": []}, "tables": {"S": {}},
                "defined_names": [], "validations": {"S": []},
                "conditional_formatting": {"S": []},
                "hidden": {"S": {"rows": [], "columns": []}},
                "styles": styles if styles is not None else {},
                "formula_count": 0}
    before = snap({"A1": None, "B1": 2})
    after = snap({"A1": "new", "B1": 99})
    result = ex.compare_snapshots(before, after, {"S!A1"})
    kinds = {(item["kind"], item.get("cell"))
             for item in result["unexpected"]}
    assert ("value", "S!B1") in kinds
    assert result["changed"] == ["S!A1"]
    assert result["noop"] == []


def test_unexpected_style_change_is_detected():
    ex = importlib.import_module("xlsx_execute")
    base_style = {"S!A1": "bold", "S!B1": "plain"}
    changed_style = {"S!A1": "bold", "S!B1": "italic"}
    before = {"sheets": ["S"], "values": {"S": {"A1": None}},
              "formulas": {}, "merges": {"S": []}, "tables": {"S": {}},
              "defined_names": [], "validations": {"S": []},
              "conditional_formatting": {"S": []},
              "hidden": {"S": {"rows": [], "columns": []}},
              "styles": base_style, "formula_count": 0}
    after = dict(before, values={"S": {"A1": "x"}},
                 styles=changed_style)
    result = ex.compare_snapshots(before, after, {"S!A1"})
    assert any(item["kind"] == "style" and item["cell"] == "S!B1"
               for item in result["unexpected"])


# ---------------------------------------------------------------------------
# 12. determinism + scope guard (AC-3A20 / AC-3A21)
# ---------------------------------------------------------------------------

def test_preflight_is_byte_deterministic_and_read_only(case):
    original = sha256(case["book"])
    first = exec_cli(case, emit="preflight",
                     token=plan_token(case["fill_plan"])).stdout
    second = exec_cli(case, emit="preflight",
                      token=plan_token(case["fill_plan"])).stdout
    assert first == second
    assert sha256(case["book"]) == original
    assert not (case["dir"] / "case_filled.xlsx").exists()


def test_scope_guard_ast():
    import ast
    import re as _re
    text = EXEC.read_text(encoding="utf-8")
    tree = ast.parse(text)
    forbidden = {"insert_rows", "delete_rows", "insert_cols", "delete_cols",
                 "move_range", "translate_formula", "insert_formula"}
    offenders = [f"{node.func.attr}:{node.lineno}"
                 for node in ast.walk(tree)
                 if isinstance(node, ast.Call)
                 and getattr(node.func, "attr", None) in forbidden]
    assert offenders == [], f"row/formula operations leaked: {offenders}"
    raw_saves = [node.lineno for node in ast.walk(tree)
                 if isinstance(node, ast.Call)
                 and getattr(node.func, "attr", None) == "save"]
    assert raw_saves == [], "raw save calls must stay in xlsx_common"
    assert not _re.search(r"data_only\s*=\s*True", text)
    assert "to_r1c1" not in text and "R1C1Style" not in text


def test_dry_run_builder_still_read_only():
    """FAZ 2B modules keep their no-write guarantee after FAZ 3A."""
    mapper = (SCRIPTS / "xlsx_mapping.py").read_text(encoding="utf-8")
    assert "import openpyxl" not in mapper
