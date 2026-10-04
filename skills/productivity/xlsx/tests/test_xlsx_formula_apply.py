#!/usr/bin/env python3
"""FAZ 4 / step 6b tests -- approved synthesis plan execution (real chain).

Runs the FULL pipeline on a real fixture workbook: parse -> family ->
synthesis plan -> `xlsx_execute.execute_synthesis_plan` (Faz 1/3C safe-write:
staging -> pre-commit QA -> backup -> atomic replace -> post-commit verify).

Guarantee coverage:
  AC-4.21 a failed execution leaves partial commit = 0 and the original
          file byte-identical (SHA-256 before == after);
  AC-4.22 values_recalculated stays False everywhere;
  AC-4.24 pre/post QA report: formula count delta == planned count and no
          unrelated formula text changed;
  D7-A   synthesis logic stays out of the executor (the executor only
         applies validated texts) and unknown plan types are rejected at
         admission, not mid-write.
"""
from __future__ import annotations

import hashlib
import sys
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))

import xlsx_common as common      # noqa: E402
import xlsx_execute as execute    # noqa: E402
import xlsx_formula_family as fam  # noqa: E402
import xlsx_formula_synth as synth  # noqa: E402


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def build_fixture(path: Path) -> Path:
    """Real xlsx: copy-down family H10:H12, one unrelated formula, H13 empty."""
    wb = Workbook()
    ws = wb.active
    ws.title = "CAP"
    for row in range(10, 13):
        ws[f"F{row}"] = 2
        ws[f"G{row}"] = 3
        ws[f"H{row}"] = f"=F{row}*G{row}"
    ws["J10"] = "=1+1"                        # unrelated formula
    wb.save(path)
    return path


def plan_for(path: Path, target=("CAP", "H13")):
    wb = load_workbook(path, data_only=False)
    cells = {c.coordinate: c.value for row in wb["CAP"].iter_rows()
             for c in row if isinstance(c.value, str) and c.value.startswith("=")}
    wb.close()
    entries = fam.build_entries(cells, "CAP")
    report = fam.discover(entries)
    texts = {entry["cell"]: entry["result"]["text"] for entry in entries}
    families_with_text = {f["family_id"]: texts[f["representative"]["cell"]]
                          for f in report["families"]}
    return synth.plan_synthesis(
        report, [{"sheet": target[0], "cell": target[1]}],
        families_with_text=families_with_text,
        target_states={target: {"has_formula": False, "has_value": False}})


# --- happy path ------------------------------------------------------------

def test_happy_path_writes_planned_formula_through_safe_chain(tmp_path):
    src = build_fixture(tmp_path / "src.xlsx")
    dst = tmp_path / "out.xlsx"
    before_sha = sha256(src)
    plan = plan_for(src)
    assert plan["generated_formula_count"] == 1

    report = execute.execute_synthesis_plan(plan, target_path=src,
                                            out_path=dst)

    assert report["committed"] is True and report["write_count"] == 1
    assert report["values_recalculated"] is False     # AC-4.22
    assert report["unexpected_changes"] == []         # AC-4.24
    out = load_workbook(dst, data_only=False)
    assert out["CAP"]["H13"].value == plan["actions"][0]["formula_text"]
    assert out["CAP"]["H13"].value == "=F13*G13"      # AC-4.12 through execution
    assert out["CAP"]["H10"].value == "=F10*G10"      # untouched source
    assert out["CAP"]["J10"].value == "=1+1"          # unrelated: unchanged
    out.close()
    assert sha256(src) == before_sha                  # source never touched


def test_in_place_creates_backup_and_changes_sha(tmp_path):
    src = build_fixture(tmp_path / "src.xlsx")
    before_sha = sha256(src)
    plan = plan_for(src)
    report = execute.execute_synthesis_plan(
        plan, target_path=src, in_place=True, approve_token=plan["plan_id"])
    assert report["committed"] is True
    backup = Path(report["backup"])
    assert backup.exists() and sha256(backup) == before_sha
    assert sha256(src) != before_sha
    out = load_workbook(src, data_only=False)
    assert out["CAP"]["H13"].value == "=F13*G13"
    out.close()


def test_in_place_without_approval_token_is_refused(tmp_path):
    src = build_fixture(tmp_path / "src.xlsx")
    before = sha256(src)
    plan = plan_for(src)
    with pytest.raises(common.XlsxError) as excinfo:
        execute.execute_synthesis_plan(plan, target_path=src, in_place=True)
    assert excinfo.value.code == "APPROVAL_REQUIRED"
    assert sha256(src) == before


def test_qa_reports_formula_count_delta_expected(tmp_path):
    src = build_fixture(tmp_path / "src.xlsx")
    plan = plan_for(src)
    report = execute.execute_synthesis_plan(plan, target_path=src,
                                            out_path=tmp_path / "o.xlsx")
    assert plan["qa_expectations"]["formula_count_delta_expected"] == 1
    assert report["write_count"] == plan["qa_expectations"][
        "formula_count_delta_expected"]


# --- failure atomicity (AC-4.21) -------------------------------------------

def test_precheck_failure_leaves_original_byte_identical(tmp_path):
    src = build_fixture(tmp_path / "src.xlsx")
    dst = tmp_path / "out.xlsx"
    before = sha256(src)
    plan = plan_for(src)
    plan["actions"].append({
        "target": {"sheet": "NOPE", "cell": "A1"}, "family_id": "F-X",
        "formula_text": "=1", "drow": 0})
    plan["generated_formula_count"] = 2
    with pytest.raises(common.XlsxError) as excinfo:
        execute.execute_synthesis_plan(plan, target_path=src, out_path=dst)
    ctx = excinfo.value.context or {}
    assert ctx["committed"] is False and ctx["write_count"] == 0
    assert ctx["values_recalculated"] is False        # AC-4.22 even on failure
    assert sha256(src) == before                      # AC-4.21
    assert not dst.exists()                           # partial commit = 0


def test_qa_refusal_leaves_original_byte_identical(tmp_path, monkeypatch):
    src = build_fixture(tmp_path / "src.xlsx")
    dst = tmp_path / "out.xlsx"
    before = sha256(src)
    plan = plan_for(src)

    def forced_qa(before_snap, after_snap, planned):
        return {"unexpected": [{"kind": "value", "cell": "CAP!ZZZ9",
                                "before": None, "after": "intruder"}],
                "formula_count_before": 2, "formula_count_after": 3}

    monkeypatch.setattr(execute, "compare_snapshots", forced_qa)
    with pytest.raises(common.XlsxError) as excinfo:
        execute.execute_synthesis_plan(plan, target_path=src, out_path=dst)
    ctx = excinfo.value.context or {}
    assert ctx["committed"] is False and ctx["write_count"] == 0
    assert sha256(src) == before                      # AC-4.21
    assert not dst.exists()


def test_target_became_formula_is_refused(tmp_path):
    src = build_fixture(tmp_path / "src.xlsx")
    plan = plan_for(src)
    wb = load_workbook(src)
    wb["CAP"]["H13"] = "=99"
    wb.save(src)
    wb.close()
    before = sha256(src)
    with pytest.raises(common.XlsxError) as excinfo:
        execute.execute_synthesis_plan(plan, target_path=src,
                                       out_path=tmp_path / "o.xlsx")
    assert excinfo.value.code == "FORMULA_MODIFICATION_FORBIDDEN"
    assert sha256(src) == before


# --- admission guards (D7-A, validated lesson) -----------------------------

def test_unknown_plan_type_is_rejected_at_admission(tmp_path):
    src = build_fixture(tmp_path / "src.xlsx")
    plan = plan_for(src)
    plan["executor_contract"]["plan_type"] = "fill"    # wrong plan type
    with pytest.raises(common.XlsxError) as excinfo:
        execute.execute_synthesis_plan(plan, target_path=src,
                                       out_path=tmp_path / "o.xlsx")
    assert excinfo.value.code == "PLAN_INVALID"
    ctx = excinfo.value.context or {}
    assert ctx.get("stage") == "admission"     # died before any write stage
    assert ctx.get("committed") is False and ctx.get("write_count") == 0


def test_plan_count_mismatch_is_rejected(tmp_path):
    src = build_fixture(tmp_path / "src.xlsx")
    plan = plan_for(src)
    plan["generated_formula_count"] = 5
    with pytest.raises(common.XlsxError) as excinfo:
        execute.execute_synthesis_plan(plan, target_path=src,
                                       out_path=tmp_path / "o.xlsx")
    assert excinfo.value.code == "PLAN_INVALID"


def test_wrong_write_policy_is_rejected(tmp_path):
    src = build_fixture(tmp_path / "src.xlsx")
    plan = plan_for(src)
    plan["write_policy"] = "write_values"
    with pytest.raises(common.XlsxError) as excinfo:
        execute.execute_synthesis_plan(plan, target_path=src,
                                       out_path=tmp_path / "o.xlsx")
    assert excinfo.value.code == "PLAN_INVALID"


def test_executor_source_contains_no_synthesis_logic():   # D7-A guard
    source = Path(execute.__file__).read_text(encoding="utf-8")
    assert "def translate_formula" not in source
    assert "INTENT_CLASSES" not in source
    assert "import xlsx_formula_synth" not in source
    assert "from xlsx_formula_synth" not in source