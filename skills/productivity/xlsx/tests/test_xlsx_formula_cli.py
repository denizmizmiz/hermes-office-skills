#!/usr/bin/env python3
"""FAZ 4 / step 7b tests -- the xlsx_formula CLI (spec sections 31-33).

End-to-end through subprocess (the way the skill invokes the CLI):
  analyze -> read-only intelligence JSON (deterministic, file untouched);
  plan    -> synthesis plan with the exact expected formula;
  execute -> out-file write through the safe chain, and in-place refused
             without `--approve <plan_id>` (APPROVAL_REQUIRED) with the
             original file SHA-256 unchanged.
"""
from __future__ import annotations

import hashlib
import json
import subprocess
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parent.parent
SCRIPTS = ROOT / "scripts"
CLI = SCRIPTS / "xlsx_formula.py"


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def build_fixture(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "CAP"
    for row in range(10, 13):
        ws[f"F{row}"] = 2
        ws[f"G{row}"] = 3
        ws[f"H{row}"] = f"=F{row}*G{row}"
    ws["J10"] = "=1+1"
    wb.save(path)
    return path


def run_cli(*args):
    proc = subprocess.run([sys.executable, str(CLI), *args],
                          capture_output=True, text=True, encoding="utf-8")
    return proc


def test_analyze_is_read_only_and_deterministic(tmp_path):
    book = build_fixture(tmp_path / "src.xlsx")
    before = sha256(book)
    first = run_cli("analyze", "--workbook", str(book))
    second = run_cli("analyze", "--workbook", str(book))
    assert first.returncode == 0, first.stderr
    assert first.stdout == second.stdout               # byte-deterministic
    bundle = json.loads(first.stdout)
    assert bundle["formulas"] == 4
    assert bundle["families"] == 1 and bundle["singles"] == 1
    assert bundle["window_breakdown"] == {"single_cell": 1}
    assert bundle["intent_breakdown"] == {"calculation": 2}
    assert bundle["values_recalculated"] is False       # never claimed
    assert sha256(book) == before                       # read-only


def test_plan_emits_the_exact_expected_formula(tmp_path):
    book = build_fixture(tmp_path / "src.xlsx")
    plan_path = tmp_path / "plan.json"
    proc = run_cli("plan", "--workbook", str(book),
                   "--targets", "CAP!H13", "--out", str(plan_path))
    assert proc.returncode == 0, proc.stderr
    plan = json.loads(plan_path.read_text(encoding="utf-8"))
    assert plan["generated_formula_count"] == 1
    assert plan["actions"][0]["formula_text"] == "=F13*G13"
    assert plan["write_policy"] == "write_formula"
    summary = json.loads(proc.stdout)
    assert summary["plan_id"] == plan["plan_id"]


def test_execute_out_mode_writes_through_the_safe_chain(tmp_path):
    book = build_fixture(tmp_path / "src.xlsx")
    out = tmp_path / "out.xlsx"
    before = sha256(book)
    plan_path = tmp_path / "plan.json"
    run_cli("plan", "--workbook", str(book), "--targets", "CAP!H13",
            "--out", str(plan_path))
    proc = run_cli("execute", "--plan", str(plan_path), "--workbook",
                   str(book), "--out", str(out))
    assert proc.returncode == 0, proc.stderr
    report = json.loads(proc.stdout)
    assert report["committed"] is True and report["write_count"] == 1
    assert report["values_recalculated"] is False
    assert report["formula_count_delta"] == 1
    wb = load_workbook(out, data_only=False)
    assert wb["CAP"]["H13"].value == "=F13*G13"
    assert wb["CAP"]["J10"].value == "=1+1"
    wb.close()
    assert sha256(book) == before


def test_execute_in_place_requires_the_approval_token(tmp_path):
    book = build_fixture(tmp_path / "src.xlsx")
    before = sha256(book)
    plan_path = tmp_path / "plan.json"
    run_cli("plan", "--workbook", str(book), "--targets", "CAP!H13",
            "--out", str(plan_path))
    denied = run_cli("execute", "--plan", str(plan_path), "--workbook",
                     str(book), "--in-place")
    assert denied.returncode == 2
    err = json.loads(denied.stderr)
    assert err["error"] == "APPROVAL_REQUIRED"
    assert sha256(book) == before                        # nothing happened

    plan = json.loads(plan_path.read_text(encoding="utf-8"))
    granted = run_cli("execute", "--plan", str(plan_path), "--workbook",
                      str(book), "--in-place",
                      "--approve", plan["plan_id"])
    assert granted.returncode == 0, granted.stderr
    report = json.loads(granted.stdout)
    assert report["committed"] is True and report["backup"]
    assert sha256(book) != before


def test_plan_exit_code_three_when_everything_is_blocked(tmp_path):
    book = build_fixture(tmp_path / "src.xlsx")
    proc = run_cli("plan", "--workbook", str(book),
                   "--targets", "CAP!H20", "--family", "NOPE")
    assert proc.returncode == 3
    summary = json.loads(proc.stdout)
    assert summary["generated_formula_count"] == 0
    assert summary["blocked"][0][1] == "NO_SOURCE_FAMILY"