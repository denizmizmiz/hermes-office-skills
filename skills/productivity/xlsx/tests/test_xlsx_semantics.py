# -*- coding: utf-8 -*-
"""FAZ 2B -- xlsx_semantics.py contract, determinism and matching tests."""
import ast
import hashlib
import json
import os
import shutil
import subprocess
import sys
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

SCRIPTS = Path(__file__).resolve().parent.parent / "scripts"
SKILL = SCRIPTS.parent
SEM = SCRIPTS / "xlsx_semantics.py"
UNDERSTAND = SCRIPTS / "xlsx_understand.py"
MAPPER = SCRIPTS / "xlsx_mapping.py"
DICTIONARY = SKILL / "references" / "semantic-dictionary.tr-en.json"
LABELS = Path(__file__).resolve().parent / "fixtures" / "labels"


def run_sem(*args, expect_ok=True):
    proc = subprocess.run([sys.executable, str(SEM), *map(str, args)],
                          capture_output=True, text=True, encoding="utf-8",
                          timeout=300)
    if expect_ok:
        assert proc.returncode == 0, proc.stderr[-2000:]
    return proc


def run_map(*args, expect_ok=True):
    proc = subprocess.run([sys.executable, str(MAPPER), *map(str, args)],
                          capture_output=True, text=True, encoding="utf-8",
                          timeout=300)
    if expect_ok:
        assert proc.returncode == 0, proc.stderr[-2000:]
    return proc


def parse(proc):
    return json.loads(proc.stdout)


def sha(path):
    hasher = hashlib.sha256()
    with open(path, "rb") as handle:
        for chunk in iter(lambda: handle.read(1 << 20), b""):
            hasher.update(chunk)
    return hasher.hexdigest()


def make_book(path, variant="clean"):
    """Deterministic fixture: header + records + formula + styled input slots."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Satis"
    headers = ["Ad", "Adet", "Birim Fiyat", "Tutar", "Notlar", "Aciklama"]
    if variant == "english":
        headers = ["Name", "Quantity", "Unit Price", "Total", "Notes",
                   "Comment"]
    if variant == "reordered":
        headers = ["Adet", "Ad", "Birim Fiyat", "Tutar", "Notlar",
                   "Aciklama"]
    for index, header in enumerate(headers, start=1):
        ws.cell(row=1, column=index, value=header)
    records = [("Kalem A", 5, 1250.5), ("Kalem B", 3, 999.9),
               ("Kalem C", 7, 100.0)]
    for offset, (name, count, price) in enumerate(records, start=2):
        if variant == "reordered":
            ws.cell(row=offset, column=1, value=count)
            ws.cell(row=offset, column=2, value=name)
        else:
            ws.cell(row=offset, column=1, value=name)
            ws.cell(row=offset, column=2, value=count)
        ws.cell(row=offset, column=3, value=price)
        if variant == "mismatch":
            ws.cell(row=offset, column=4, value="yok")
        else:
            ws.cell(row=offset, column=4, value=f"=B{offset}*C{offset}")
    for row in range(2, 5):
        ws.cell(row=row, column=3).number_format = "#,##0.00 ₺"
    # evidenced input slots: styled empty cells (fill) with a list validation
    fill = PatternFill("solid", fgColor="FFF2CC")
    for row in range(6, 9):
        for column in ("E", "F"):
            cell = ws[f"{column}{row}"]
            cell.fill = fill
            cell.number_format = "General"
    ws["E5"] = "Notlar"
    validation = DataValidation(type="list", formula1='"A,B,C"',
                                allow_blank=True)
    ws.add_data_validation(validation)
    validation.add("E6:E8")
    if variant == "extra":
        ws["G1"] = "Fazla Kolon"
        ws["G2"] = "x"
        ws["G3"] = "y"
    if variant == "missing":
        ws["D1"] = "Bilinmeyen Alan"
        for row in range(2, 5):
            # NOTE: ws.cell(value=None) is a NO-OP in openpyxl; assign the
            # attribute so the cell value is really cleared.
            ws.cell(row=row, column=4).value = None
    if variant == "duplicate":
        ws["C1"] = "Ad"
        ws["C2"] = "Kalem Z"
    if variant == "ambiguous":
        # two columns that both look exactly like the template's "Birim Fiyat"
        ws["C1"] = "Birim Fiyat"
        ws["D1"] = "Birim Fiyat"
        for row in range(2, 5):
            cell = ws.cell(row=row, column=4)
            cell.value = float(row * 10)
            cell.number_format = "#,##0.00 ₺"
    if variant == "hidden":
        ws.row_dimensions[3].hidden = True
    if variant == "merged":
        # regression: 2A reports merges as dicts, profiles as strings
        ws.merge_cells("E7:F7")
    if variant == "lookup":
        lookup = DataValidation(type="list", formula1="$H$1:$H$5",
                                allow_blank=True)
        ws.add_data_validation(lookup)
        lookup.add("B2:B4")
    wb.save(path)
    return path


@pytest.fixture(scope="module")
def books(tmp_path_factory):
    directory = tmp_path_factory.mktemp("fx2b")
    paths = {}
    for name in ("clean", "english", "reordered", "mismatch", "extra",
                 "missing", "duplicate", "hidden", "lookup", "ambiguous",
                 "merged"):
        paths[name] = make_book(directory / f"{name}.xlsx", variant=name)
    return paths


# ---------------------------------------------------------------------------
# contract
# ---------------------------------------------------------------------------

def test_semantics_contract(books):
    doc = parse(run_sem(books["clean"]))
    for key in ("ok", "mode", "command", "warnings", "unsupported",
                "diagnostics", "semantics", "identity"):
        assert key in doc, key
    assert doc["ok"] is True and doc["command"] == "semantics"
    sheet = doc["semantics"]["sheets"][0]
    assert sheet["name"] == "Satis"
    assert doc["semantics"]["dictionary_version"] == "tr-en.1"


def test_missing_input_is_structured(tmp_path):
    proc = run_sem(tmp_path / "nope.xlsx", expect_ok=False)
    assert proc.returncode != 0
    error = json.loads(proc.stderr)
    assert error["ok"] is False
    assert error["error_code"] == "FILE_NOT_FOUND"
    assert error["recovery"]


def test_bad_doc_map_is_structured(tmp_path):
    bad = tmp_path / "bad.json"
    bad.write_text("{ not json", encoding="utf-8")
    proc = run_sem("--doc-map", bad, expect_ok=False)
    assert proc.returncode != 0
    assert json.loads(proc.stderr)["error_code"] == "VALIDATION_FAILED"
    shape = tmp_path / "shape.json"
    shape.write_text('{"ok": true}', encoding="utf-8")
    proc = run_sem("--doc-map", shape, expect_ok=False)
    assert json.loads(proc.stderr)["error_code"] == "VALIDATION_FAILED"


def test_conflicting_arguments(books, tmp_path):
    profile = tmp_path / "p.json"
    profile.write_text("{}", encoding="utf-8")
    proc = run_sem(books["clean"], "--emit", "profile", "--match", profile,
                   expect_ok=False)
    assert proc.returncode != 0
    assert json.loads(proc.stderr)["error_code"] == "SPEC_INVALID"


# ---------------------------------------------------------------------------
# B1 -- roles, evidence, coverage
# ---------------------------------------------------------------------------

def test_roles_have_scores_and_evidence(books):
    doc = parse(run_sem(books["clean"]))
    sheet = doc["semantics"]["sheets"][0]
    assert sheet["coverage"]["columns"] == len(sheet["columns"]) > 0
    for column in sheet["columns"]:
        assert column["semantic_role"]
        assert column["value_kind"] in (
            "string", "integer", "decimal", "date", "datetime", "boolean",
            "formula", "blank", "mixed")
        assert 0 <= column["role_score"] <= 100
        assert column["role_level"] in ("HIGH", "MEDIUM", "LOW")
        assert column["evidence"], column["letter"]
        assert column["role_score"] == column["confidence_score"]
        assert column["role_level"] == column["confidence_level"]
        if column["requires_review"]:
            assert column["ambiguity"]
    assert sheet["coverage"]["unknown"] == 0


def test_expected_roles_on_clean_book(books):
    doc = parse(run_sem(books["clean"]))
    roles = {c["header"]: c["semantic_role"]
             for c in doc["semantics"]["sheets"][0]["columns"]}
    assert roles["Adet"] == "quantity"
    assert roles["Birim Fiyat"] == "currency"
    assert roles["Tutar"] == "currency"


def test_low_confidence_is_never_silent(books):
    doc = parse(run_sem(books["english"]))
    low_columns = [c for c in doc["semantics"]["sheets"][0]["columns"]
                   if c["role_level"] == "LOW"]
    codes = [w["code"] for w in doc["warnings"]]
    for _ in low_columns:
        assert "LOW_CONFIDENCE_DETECTION" in codes


def test_value_kind_is_not_the_role(books):
    doc = parse(run_sem(books["clean"]))
    by_header = {c["header"]: c for c in doc["semantics"]["sheets"][0]["columns"]}
    assert by_header["Tutar"]["semantic_role"] == "currency"
    assert by_header["Tutar"]["value_kind"] == "formula"


def test_unit_guess_requires_evidence(books):
    doc = parse(run_sem(books["clean"]))
    by_header = {c["header"]: c for c in doc["semantics"]["sheets"][0]["columns"]}
    assert by_header["Birim Fiyat"]["unit_guess"] == "₺"
    assert by_header["Birim Fiyat"]["unit_evidence"]
    assert by_header["Ad"]["unit_guess"] is None
    assert by_header["Ad"]["unit_evidence"] == []


def test_dictionary_is_valid_and_unique(books):
    data = json.loads(DICTIONARY.read_text(encoding="utf-8"))
    assert data["dictionary_version"] == "tr-en.1"
    seen = {}
    for role, terms in data["roles"].items():
        assert terms == sorted(terms, key=str.casefold)
        for term in terms:
            key = term.casefold().strip()
            assert key not in seen, f"{term} in {role} and {seen[key]}"
            seen[key] = role
    assert data["units"]
    proc = subprocess.run(
        [sys.executable, "-c",
         "import sys; sys.path.insert(0, r'%s'); import xlsx_semantics as s;"
         "d = s.load_dictionary(); print(len(d['conflicts']))" % SCRIPTS],
        capture_output=True, text=True, encoding="utf-8")
    assert proc.stdout.strip() == "0"


def test_duplicate_headers_create_ambiguity(books):
    doc = parse(run_sem(books["duplicate"]))
    sheet = doc["semantics"]["sheets"][0]
    ambiguous = [c for c in sheet["columns"] if c["ambiguity"]]
    assert ambiguous, "duplicate/identical headers must surface ambiguity"
    for column in ambiguous:
        assert column["requires_review"] is True
        entry = column["ambiguity"][0]
        assert entry["margin"] < 10
        assert entry["policy"]


# ---------------------------------------------------------------------------
# B2/B3 -- blocks, slots, profile identity
# ---------------------------------------------------------------------------

def test_blocks_and_slots(books, tmp_path):
    out = tmp_path / "p.json"
    doc = parse(run_sem(books["clean"], "--emit", "profile", "--out", out))
    profile = doc["profile"]
    sheet = profile["sheets"][0]
    archetypes = {b["archetype"] for b in sheet["blocks"]}
    assert "record_block" in archetypes
    record = next(b for b in sheet["blocks"] if b["archetype"] == "record_block")
    assert record["rows"][0] >= 2
    assert record["record_key_candidate"] is not None
    assert record["confidence_level"] in ("HIGH", "MEDIUM", "LOW")
    assert record["evidence"]
    assert sheet["slots"], "styled empty input slot must be reported"
    for slot in sheet["slots"]:
        assert slot["required_source"] == "heuristic"
        assert slot["evidence"]
    for column in sheet["columns"]:
        assert column["required_source"] == "heuristic"
        assert column["required_evidence"]


def test_profile_determinism(books, tmp_path):
    same_args = run_sem(books["clean"], "--emit", "profile")
    assert same_args.stdout == run_sem(books["clean"], "--emit",
                                       "profile").stdout
    first = run_sem(books["clean"], "--emit", "profile",
                    "--out", tmp_path / "a.json")
    second = run_sem(books["clean"], "--emit", "profile",
                     "--out", tmp_path / "b.json")
    assert (tmp_path / "a.json").read_bytes() == (tmp_path / "b.json").read_bytes()
    assert parse(first)["profile"] == parse(second)["profile"]
    assert parse(first)["profile_id"] == parse(second)["profile_id"]


def test_profile_is_mtime_independent(books, tmp_path):
    source = books["clean"]
    left_dir = tmp_path / "left"
    right_dir = tmp_path / "right"
    left_dir.mkdir()
    right_dir.mkdir()
    left = left_dir / "same.xlsx"
    right = right_dir / "same.xlsx"
    shutil.copy(source, left)
    shutil.copy(source, right)
    os.utime(left, (1_600_000_000, 1_600_000_000))
    os.utime(right, (1_700_000_000, 1_700_000_000))
    a = run_sem(left, "--emit", "profile", "--out", left_dir / "p.json")
    b = run_sem(right, "--emit", "profile", "--out", right_dir / "p.json")
    assert (left_dir / "p.json").read_bytes() == (right_dir / "p.json").read_bytes()
    assert parse(a)["profile_id"] == parse(b)["profile_id"]
    assert parse(a)["profile"]["content_identity"] == \
        parse(b)["profile"]["content_identity"]
    assert parse(a)["profile"]["source_bytes_sha256"] == \
        parse(b)["profile"]["source_bytes_sha256"]


def test_profile_identity_follows_content(books, tmp_path):
    base = parse(run_sem(books["clean"], "--emit", "profile"))["profile_id"]
    extra = parse(run_sem(books["extra"], "--emit", "profile"))["profile_id"]
    assert base != extra
    again = parse(run_sem(books["clean"], "--emit", "profile"))["profile_id"]
    assert base == again


def test_doc_map_input_matches_workbook_input(books, tmp_path):
    fresh = parse(run_sem(books["clean"]))
    map_file = tmp_path / "map.json"
    proc = subprocess.run([sys.executable, str(UNDERSTAND),
                           str(books["clean"])], capture_output=True,
                          text=True, encoding="utf-8", timeout=300)
    assert proc.returncode == 0, proc.stderr[-500:]
    map_file.write_text(proc.stdout, encoding="utf-8")
    replayed = parse(run_sem("--doc-map", map_file))
    assert replayed["identity"]["profile_id"] == fresh["identity"]["profile_id"]
    assert replayed["semantics"] == fresh["semantics"]


# ---------------------------------------------------------------------------
# B4 -- matching
# ---------------------------------------------------------------------------

def _profile_file(books, tmp_path, name="clean"):
    out = tmp_path / f"{name}.profile.json"
    run_sem(books[name], "--emit", "profile", "--out", out)
    return out


def test_self_match_is_high(books, tmp_path):
    profile = _profile_file(books, tmp_path)
    report = parse(run_sem(books["clean"], "--match", profile))["match_report"]
    assert report["overall_score"] >= 80
    assert report["overall_level"] == "HIGH"
    assert report["same_content_identity"] is True
    for key in ("region_score", "column_score", "formula_score",
                "structural_score"):
        assert key in report["components"]
    assert report["unmatched_template"] == []
    assert report["extra_instance"] == []
    assert report["conflicts"] == []


def test_variant_missing_column(books, tmp_path):
    profile = _profile_file(books, tmp_path)
    report = parse(run_sem(books["missing"], "--match", profile))["match_report"]
    headers = [item.get("header") for item in report["unmatched_template"]]
    assert "Tutar" in headers
    assert report["overall_score"] < 100


def test_variant_extra_column(books, tmp_path):
    profile = _profile_file(books, tmp_path)
    report = parse(run_sem(books["extra"], "--match", profile))["match_report"]
    headers = [item.get("header") for item in report["extra_instance"]]
    assert "Fazla Kolon" in headers


def test_variant_type_mismatch_is_a_conflict(books, tmp_path):
    profile = _profile_file(books, tmp_path)
    report = parse(run_sem(books["mismatch"], "--match", profile))["match_report"]
    kinds = {item["kind"] for item in report["conflicts"]}
    assert "value_kind" in kinds


def test_variant_reordered_still_matches(books, tmp_path):
    profile = _profile_file(books, tmp_path)
    report = parse(run_sem(books["reordered"], "--match", profile))["match_report"]
    assert report["overall_score"] >= 65
    assert report["components"]["column_score"] == 100


def test_match_ambiguity_is_flagged(books, tmp_path):
    profile = _profile_file(books, tmp_path)
    report = parse(run_sem(books["ambiguous"], "--match", profile))["match_report"]
    assert report["requires_review"] is True
    ambiguous = [entry for sheet in report["sheets"]
                 for entry in sheet["column_matches"] if entry["ambiguity"]]
    assert ambiguous
    assert any(entry["requires_review"] for entry in ambiguous)


def test_match_is_deterministic(books, tmp_path):
    profile = _profile_file(books, tmp_path)
    first = run_sem(books["clean"], "--match", profile)
    second = run_sem(books["clean"], "--match", profile)
    assert first.stdout == second.stdout


# ---------------------------------------------------------------------------
# safety: AST guard + SHA no-write
# ---------------------------------------------------------------------------

FORBIDDEN_ATTRS = {"save", "save_workbook_safe", "atomic_replace",
                   "backup_workbook"}


def _violations(source_path):
    tree = ast.parse(Path(source_path).read_text(encoding="utf-8"))
    found = []
    for node in ast.walk(tree):
        if isinstance(node, ast.Call):
            func = node.func
            name = getattr(func, "attr", None) or getattr(func, "id", None)
            if name in FORBIDDEN_ATTRS:
                found.append(f"{name}() at line {node.lineno}")
            if name == "ZipFile":
                found.append(f"ZipFile() at line {node.lineno}")
            if name == "open":
                for arg in node.args[1:]:
                    if isinstance(arg, ast.Constant) and isinstance(arg.value, str) \
                            and (("w" in arg.value) or ("a" in arg.value)):
                        found.append(f"open(mode={arg.value!r}) at line "
                                     f"{node.lineno}")
    return found


def test_no_write_ast_guard():
    for module in (SEM, MAPPER):
        text = Path(module).read_text(encoding="utf-8")
        assert "import openpyxl" not in text, module.name
        assert _violations(module) == [], f"{module.name}: {_violations(module)}"


def test_real_files_are_untouched(books):
    before = {name: sha(path) for name, path in books.items()}
    run_sem(books["clean"])
    run_sem(books["hidden"], "--emit", "profile")
    run_sem(books["lookup"], "--match", _profile_output(books))
    after = {name: sha(path) for name, path in books.items()}
    assert before == after


def _profile_output(books):
    proc = subprocess.run([sys.executable, "-c",
                           "import subprocess, sys, tempfile, pathlib;"
                           "import json;"
                           "tmp = pathlib.Path(tempfile.mkdtemp())/'p.json';"
                           "subprocess.run([sys.executable, r'%s', r'%s',"
                           "'--emit','profile','--out',str(tmp)], check=True,"
                           "capture_output=True);"
                           "print(tmp)" % (SEM, books["clean"])],
                          capture_output=True, text=True, encoding="utf-8")
    return proc.stdout.strip()


def test_cap_list_reports_truncation():
    sys.path.insert(0, str(SCRIPTS))
    import xlsx_common as common  # noqa: PLC0415
    import xlsx_semantics as semantics  # noqa: PLC0415
    result = common.Result()
    listed, meta = semantics.cap_list(list(range(5)), 2, result,
                                      "SLOTS_TRUNCATED", "slots")
    assert listed == [0, 1]
    assert meta == {"count_total": 5, "returned_count": 2, "truncated": True}
    assert result.warnings and result.warnings[0]["code"] == "SLOTS_TRUNCATED"
    listed, meta = semantics.cap_list([1], 5, result, "X", "y")
    assert meta["truncated"] is False


# ---------------------------------------------------------------------------
# oracle accuracy (AC-B03) -- runs only when the user's labels exist
# ---------------------------------------------------------------------------

KNOWN_VARIANTS = {"clean", "english", "reordered", "mismatch", "extra",
                 "missing", "duplicate", "hidden", "lookup", "ambiguous",
                 "merged"}


def _fixture_variant(name):
    """Fixture book name -> make_book variant.

    Regression: the oracle used to build every labeled book as "clean",
    so labels for `english`/`reordered`/... were graded against the wrong
    workbook. Unknown names still fall back to "clean".
    """
    return name if name in KNOWN_VARIANTS else "clean"


def test_role_accuracy_against_user_labels(books, tmp_path):
    label_file = LABELS / "labels.json"
    if not label_file.exists():
        pytest.skip("user-labeled oracle not provided yet "
                    "(see tests/fixtures/labels/README.md)")
    labels = json.loads(label_file.read_text(encoding="utf-8"))
    if not labels.get("columns"):
        pytest.skip("labels.json has no columns yet")
    directory = tmp_path / "oracle"
    directory.mkdir()
    target = None
    for name in labels.get("books") or {}:
        target = make_book(directory / f"{name}.xlsx",
                           variant=_fixture_variant(name))
    total, correct, low_total, low_correct = 0, 0, 0, 0
    for entry in labels["columns"]:
        book = make_book(directory / f"{entry['book']}.xlsx",
                         variant=_fixture_variant(entry["book"]))
        doc = parse(run_sem(book))
        sheet = doc["semantics"]["sheets"][0]
        column = next((c for c in sheet["columns"]
                       if c["letter"] == entry["column"]), None)
        assert column is not None, entry
        total += 1
        if column["semantic_role"] == entry["expected_role"]:
            correct += 1
            if column["role_level"] == "LOW":
                low_correct += 1
        if column["role_level"] == "LOW":
            low_total += 1
    accuracy = correct / total
    assert accuracy >= 0.90, f"role accuracy {accuracy:.2%} ({correct}/{total})"


# ---------------------------------------------------------------------------
# review_state -- the human decision next to the machine decision (delta #7)
# ---------------------------------------------------------------------------

def test_review_state_records_human_decisions(books, tmp_path):
    profile = _profile_file(books, tmp_path)
    report = parse(run_sem(books["ambiguous"], "--match",
                           profile))["match_report"]
    state = report["review_state"]
    assert state["state"] == "unreviewed"
    assert state["subjects"], "an ambiguous match must list review subjects"
    assert all(item["state"] == "unreviewed" for item in state["subjects"])

    review = tmp_path / "review.json"
    review.write_text(json.dumps({"decisions": {
        item["subject"]: {"state": "accepted", "note": "looks right"}
        for item in state["subjects"]}}), encoding="utf-8")
    second = parse(run_sem(books["ambiguous"], "--match", profile,
                           "--review-file", review))["match_report"]
    assert second["review_state"]["state"] == "reviewed"
    # a review never rewrites the machine decision
    assert second["overall_score"] == report["overall_score"]
    assert second["requires_review"] == report["requires_review"]
    first_cols = [entry["ambiguity"] for sheet in report["sheets"]
                  for entry in sheet["column_matches"]]
    second_cols = [entry["ambiguity"] for sheet in second["sheets"]
                   for entry in sheet["column_matches"]]
    assert first_cols == second_cols


def test_review_state_partial_and_unknown_subject(books, tmp_path):
    profile = _profile_file(books, tmp_path)
    pending = parse(run_sem(books["ambiguous"], "--match",
                            profile))["match_report"]
    subjects = [item["subject"] for item in pending["review_state"]["subjects"]]
    assert subjects, "an ambiguous match must list its review subjects"
    decisions = {subjects[0]: {"state": "pending"},
                 "column|Z->Z": "accepted"}
    review = tmp_path / "review.json"
    review.write_text(json.dumps({"decisions": decisions}), encoding="utf-8")
    doc = parse(run_sem(books["ambiguous"], "--match", profile,
                        "--review-file", review))
    assert doc["match_report"]["review_state"]["unknown_review_subjects"] == \
        ["column|Z->Z"]
    assert doc["match_report"]["review_state"]["state"] == "partially_reviewed"
    codes = {item["code"] for item in doc["warnings"]}
    assert "UNKNOWN_REVIEW_SUBJECT" in codes
    assert "REVIEW_INCOMPLETE" in codes


def test_review_state_is_empty_when_nothing_is_ambiguous(books, tmp_path):
    profile = _profile_file(books, tmp_path)
    report = parse(run_sem(books["clean"], "--match",
                           profile))["match_report"]
    assert report["review_state"]["state"] == "nothing_to_review"
    assert report["review_state"]["subjects"] == []


def test_review_file_requires_match(books, tmp_path):
    review = tmp_path / "review.json"
    review.write_text(json.dumps({"decisions": {}}), encoding="utf-8")
    proc = run_sem(books["clean"], "--review-file", review, expect_ok=False)
    assert json.loads(proc.stderr)["error_code"] == "SPEC_INVALID"


def test_role_coverage_threshold_on_fixtures(books):
    """AC-B02, measured against the fixtures (unknown is never a silent exit).

    Coverage is computed here from the raw columns -- not read from the
    module's own coverage counter -- so the module cannot grade itself.
    """
    for name in ("clean", "english", "reordered", "mismatch", "duplicate",
                 "hidden"):
        doc = parse(run_sem(books[name]))
        labelable = labelled = unknown = 0
        for sheet in doc["semantics"]["sheets"]:
            for col in sheet["columns"]:
                has_header = bool((col.get("header") or "").strip())
                has_data = col.get("data_type") not in (None, "blank",
                                                        "empty")
                if not (has_header or has_data):
                    continue
                labelable += 1
                if col.get("semantic_role") in (None, "unknown"):
                    unknown += 1
                else:
                    labelled += 1
        assert labelable, name
        assert labelled / labelable >= 0.95, (name, labelled, labelable)
        if unknown:
            codes = {item["code"] for item in doc["warnings"]}
            assert "LOW_CONFIDENCE_DETECTION" in codes, (name, unknown)


def test_hidden_rows_and_columns_are_supported(tmp_path):
    """Regression: 2A emits hidden columns as objects (index / outline
    level), not plain integers. Sorting them raised TypeError on real
    workbooks, so profile/matching failed on files that merely hid a
    column. Both hidden rows and hidden columns must work, and the hidden
    state is part of the content identity.
    """
    def build(path, hidden):
        wb = Workbook()
        ws = wb.active
        ws.title = "Gizli"
        ws.append(["Urun No", "Adet", "Birim Fiyat"])
        for offset, row in enumerate([("K-1", 2, 3.5), ("K-2", 5, 1.25)],
                                     start=2):
            for index, value in enumerate(row, start=1):
                ws.cell(row=offset, column=index, value=value)
        if hidden:
            ws.row_dimensions[2].hidden = True
            ws.column_dimensions["B"].hidden = True
        wb.save(path)
        return path

    hidden_book = build(tmp_path / "hidden.xlsx", True)
    plain_book = build(tmp_path / "plain.xlsx", False)

    first = run_sem(hidden_book, "--emit", "profile")
    profile = parse(first)
    assert profile["profile_id"]
    second = run_sem(hidden_book, "--emit", "profile")
    assert first.stdout == second.stdout, "profile must stay deterministic"

    hidden_id = profile["profile"]["content_identity"]
    plain_id = parse(run_sem(plain_book, "--emit", "profile"))
    assert hidden_id != plain_id["profile"]["content_identity"], \
        "hidden rows/columns are part of the semantic content identity"

    out = tmp_path / "h.profile.json"
    run_sem(hidden_book, "--emit", "profile", "--out", out)
    report = parse(run_sem(hidden_book, "--match", out))["match_report"]
    assert report["overall_score"] >= 80
    assert report["components"]["structural_score"] == 100


# ---------------------------------------------------------------------------
# caps / structure parity (2A policy: totals visible, never silent)
# ---------------------------------------------------------------------------

def _wide_book(path, columns=2100, title="Genis"):
    wb = Workbook()
    ws = wb.active
    ws.title = title
    for column in range(1, columns + 1):
        ws.cell(row=1, column=column, value="K%d" % column)
        ws.cell(row=2, column=column, value=column)
        ws.cell(row=3, column=column, value=column * 2)
    wb.save(path)
    return path


def test_wide_sheet_profile_is_capped_with_warning(tmp_path):
    """A very wide sheet caps the column list, keeps the totals and stays
    under the profile size bound (AC-B26) - never silently."""
    book = _wide_book(tmp_path / "wide.xlsx")
    out = tmp_path / "wide.profile.json"
    payload = parse(run_sem(book, "--emit", "profile", "--out", out))
    sheet = json.loads(out.read_text(encoding="utf-8"))["sheets"][0]
    cap = sheet["columns_cap"]
    assert cap["truncated"] is True
    assert cap["count_total"] == 2100
    assert cap["returned_count"] == 2000 == len(sheet["columns"])
    assert any(w["code"] == "COLUMNS_TRUNCATED" for w in payload["warnings"])
    blanks = sheet["blank_columns"]
    assert blanks["material_count"] + blanks["count"] == cap["count_total"]
    assert out.stat().st_size <= 2 * 1024 * 1024


def test_match_survives_merged_cells(books, tmp_path):
    """Regression: the Document Map reports merges as dicts while the
    profile stores range strings; matching must not crash on that."""
    profile = tmp_path / "clean.json"
    run_sem(books["clean"], "--emit", "profile", "--out", profile)
    same = parse(run_sem(books["clean"], "--match", profile))["match_report"]
    assert same["sheets"][0]["components"]["structural_score"] == 100

    other = parse(run_sem(books["merged"], "--match", profile))["match_report"]
    report = other["sheets"][0]
    # only the merge component differs: (0 + 100*5) / 6
    assert report["components"]["structural_score"] == 83
    assert any("merged ranges shared 0/1" in item
               for item in report["evidence"])
    assert report["pair_level"] in ("HIGH", "MEDIUM", "LOW")
    assert "blank_instance_columns" in report
    assert report["blank_instance_columns"]["material_count"] > 0
    assert report["blank_instance_columns"]["count"] == 0


def test_wide_instance_is_capped_before_matching(books, tmp_path):
    """Matching is O(template x instance): a 16k-column instance must be
    capped WITH a warning, and the totals must stay visible."""
    profile = tmp_path / "clean.json"
    run_sem(books["clean"], "--emit", "profile", "--out", profile)
    wide = _wide_book(tmp_path / "wide_instance.xlsx", title="Satis")
    payload = parse(run_sem(wide, "--match", profile))
    report = payload["match_report"]["sheets"][0]
    cap = report["instance_columns_cap"]
    assert cap["truncated"] is True
    assert cap["count_total"] == 2100
    assert cap["returned_count"] == 2000
    assert any(w["code"] == "INSTANCE_COLUMNS_TRUNCATED"
               for w in payload["warnings"])
    assert report["column_matches"]
