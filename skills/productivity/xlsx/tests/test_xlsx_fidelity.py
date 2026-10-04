"""Fidelity round-trip tests: the safe-write chain must preserve features.

Builds one workbook carrying the Phase-0.5 18-feature matrix (minus the
advanced Excel drawings that openpyxl cannot generate -- see the note at
the bottom), pushes it through xlsx_common's load -> save chain, and
compares every feature on the other side. Also checks the OOXML package
parts (media / chart XML) still exist.

The image check needs Pillow; when Pillow is missing that one check is
skipped and the rest still runs (report it as NOT_TESTED).
"""
from __future__ import annotations

import base64
import io
import json
import subprocess
import sys
import zipfile
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.drawing.image import Image as XLImage
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

SCRIPTS = Path(__file__).resolve().parent.parent / "scripts"

PNG_1PX = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNk"
    "YPhfDwAChwGA60e6kgAAAABJRU5ErkJggg==")


def build_fixture(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"

    # formula + font + fill + comment + hyperlink
    ws["A1"] = "Feature"
    ws["A1"].font = Font(bold=True, color="FF0000", size=14)
    ws["A2"] = "filled"
    ws["A2"].fill = PatternFill("solid", fgColor="DDEBF7")
    ws["A3"] = "commented"
    ws["A3"].comment = Comment("the note text", "the author")
    ws["A4"] = "clickable"
    ws["A4"].hyperlink = "https://example.com"
    ws["B2"] = "=SUM(B10:B12)"

    # hidden row / hidden column / freeze panes / merge
    ws.row_dimensions[5].hidden = True
    ws.column_dimensions["D"].hidden = True
    ws.freeze_panes = "A2"
    ws.merge_cells("F1:G1")

    # data rows for table / validation / CF / chart (string headers!)
    ws["A10"], ws["B10"], ws["C10"] = "h1", "h2", "h3"
    ws["A11"], ws["B11"], ws["C11"] = "x", 1, 3
    ws["A12"], ws["B12"], ws["C12"] = "y", 2, 2
    ws["A13"], ws["B13"], ws["C13"] = "z", 3, 1

    # native table
    table = Table(displayName="SalesT", ref="A10:C13")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium9",
                                          showRowStripes=True)
    ws.add_table(table)

    # data validation + conditional formatting
    dv = DataValidation(type="list", formula1='"a,b"', allow_blank=True)
    dv.add("B11:B13")
    ws.add_data_validation(dv)
    ws.conditional_formatting.add(
        "C11:C13",
        CellIsRule(operator="greaterThan", formula=["1"],
                   fill=PatternFill("solid", fgColor="C6EFCE")))

    # chart
    chart = BarChart()
    chart.title = "Fidelity"
    data = Reference(ws, min_col=2, min_row=11, max_col=2, max_row=13)
    chart.add_data(data, titles_from_data=False)
    ws.add_chart(chart, "J2")

    # image (needs Pillow to embed; fixture keeps going without it)
    try:
        img = XLImage(io.BytesIO(PNG_1PX))
        ws.add_image(img, "J20")
    except ImportError:
        pass

    # defined name + hidden sheet + sheet order
    wb.defined_names["RateName"] = DefinedName(
        "RateName", attr_text="'Data'!$B$2")
    ws2 = wb.create_sheet("Notes")
    ws2["A1"] = "notes"
    ws3 = wb.create_sheet("Gizli")
    ws3["A1"] = "hidden sheet"
    ws3.sheet_state = "hidden"

    wb.save(path)
    return path


def snapshot(path):
    """Extract the comparable feature inventory of a workbook."""
    wb = load_workbook(path, data_only=False)
    ws = wb["Data"]
    try:
        image_count = len(getattr(ws, "_images", []))
    except Exception:  # noqa: BLE001
        image_count = -1
    with zipfile.ZipFile(path) as zf:
        names = zf.namelist()
    return {
        "sheetnames": list(wb.sheetnames),
        "sheet_states": {name: s.sheet_state for name, s in
                         ((n, wb[n]) for n in wb.sheetnames)},
        "formula": ws["B2"].value,
        "font_bold": ws["A1"].font.bold,
        "font_color": ws["A1"].font.color.rgb if ws["A1"].font.color else None,
        "font_size": ws["A1"].font.size,
        "fill": ws["A2"].fill.fgColor.rgb,
        "comment_text": ws["A3"].comment.text if ws["A3"].comment else None,
        "comment_author": (ws["A3"].comment.author
                           if ws["A3"].comment else None),
        "hyperlink": ws["A4"].hyperlink.target if ws["A4"].hyperlink else None,
        "row5_hidden": bool(ws.row_dimensions[5].hidden),
        "colD_hidden": bool(ws.column_dimensions["D"].hidden),
        "freeze": ws.freeze_panes,
        "merges": sorted(str(r) for r in ws.merged_cells.ranges),
        "tables": {t.displayName: t.ref for t in ws.tables.values()},
        "defined_names": {n: dn.attr_text
                          for n, dn in wb.defined_names.items()},
        "validations": [str(dv.sqref)
                        for dv in ws.data_validations.dataValidation],
        "cf_count": len(list(ws.conditional_formatting)),
        "chart_count": len(ws._charts),
        "chart_type": type(ws._charts[0]).__name__ if ws._charts else None,
        "image_count": image_count,
        "media_parts": sorted(n for n in names if n.startswith("xl/media/")),
        "chart_parts": sorted(n for n in names
                              if n.startswith("xl/charts/")),
    }


@pytest.fixture
def fixture_path(tmp_path):
    return build_fixture(tmp_path / "fidelity.xlsx")


def test_roundtrip_preserves_feature_matrix(fixture_path, tmp_path):
    before = snapshot(fixture_path)
    out = tmp_path / "roundtrip.xlsx"

    proc = subprocess.run(
        [sys.executable, "-c",
         "import sys\n"
         f"sys.path.insert(0, r'{SCRIPTS.as_posix()}')\n"
         "import xlsx_common as c\n"
         "h = c.load_workbook_safe(sys.argv[1])\n"
         "info = c.save_workbook_safe(h, sys.argv[2], backup=False)\n"
         "print('SAVED' if info['ok'] else 'FAILED')",
         str(fixture_path), str(out)],
        capture_output=True, text=True, encoding="utf-8")
    assert proc.returncode == 0 and "SAVED" in proc.stdout, proc.stderr

    after = snapshot(out)

    # every feature of the matrix must survive the load -> save chain
    comparable = [k for k in before if k != "image_count"]
    diffs = {k: (before[k], after[k]) for k in comparable
             if before[k] != after[k]}
    assert diffs == {}, f"features changed in round-trip: {diffs}"

    # the OOXML package parts are still real parts, not stripped
    assert after["chart_parts"], "chart XML part missing after round-trip"
    if before["media_parts"]:
        assert after["media_parts"], "media part missing after round-trip"


def test_image_roundtrip_or_not_tested(fixture_path, tmp_path):
    pytest.importorskip("PIL",
                        reason="Pillow not installed: image fidelity NOT_TESTED")
    before = snapshot(fixture_path)
    assert before["image_count"] == 1 and before["media_parts"]
    out = tmp_path / "img.xlsx"
    proc = subprocess.run(
        [sys.executable, "-c",
         "import sys\n"
         f"sys.path.insert(0, r'{SCRIPTS.as_posix()}')\n"
         "import xlsx_common as c\n"
         "h = c.load_workbook_safe(sys.argv[1])\n"
         "c.save_workbook_safe(h, sys.argv[2], backup=False)",
         str(fixture_path), str(out)],
        capture_output=True, text=True, encoding="utf-8")
    assert proc.returncode == 0, proc.stderr
    after = snapshot(out)
    assert after["image_count"] == before["image_count"]
    assert after["media_parts"] == before["media_parts"]


# NOT_TESTED (report honestly, do NOT mark as preserved):
#   text box, arrow, shape, SmartArt, sparkline -- these Excel-generated
#   drawing objects cannot be produced by openpyxl 3.1.5, so the round-trip
#   behaviour of the safe-write chain for them is unmeasured in this suite.
